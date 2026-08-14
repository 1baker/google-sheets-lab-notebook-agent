from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from lab_notebook_agent.google_sheets import (
    google_scientist_workspace_upgrade_requests,
    reaction_outcomes_setup_requests,
)
from lab_notebook_agent.reaction_outcomes import (
    excel_reaction_outcome_formula,
    google_reaction_outcome_array_formulas,
)
from lab_notebook_agent.schema import sheet_by_name
from lab_notebook_agent.scientist_workspace import (
    daily_log_rows_from_tables,
    google_reaction_master_array_formulas,
    plot_studio_measurement_formula,
    plot_studio_process_formula,
    reaction_master_excel_formula,
    result_rows_from_tables,
    summarize_record_governance,
)
from lab_notebook_agent.templates import save_workbook


class ScientistWorkspaceTests(unittest.TestCase):
    def test_compact_entry_rows_project_into_canonical_science_tables(self) -> None:
        tables = {
            "Daily Log": [],
            "Bench Log": [
                {
                    "Run ID": "EP-900",
                    "Date & time": "2026-08-12T10:00:00",
                    "Stage": "feed",
                    "Temperature (°C)": 70,
                    "RPM": 300,
                    "Observation / action": "Feed started cleanly.",
                }
            ],
            "Results": [],
            "Measurements": [
                {
                    "Run ID": "EP-900",
                    "Sample ID": "EP-900-S1",
                    "Measurement": "DLS particle size",
                    "Numeric value": 185,
                    "Units": "nm",
                }
            ],
        }
        log = daily_log_rows_from_tables(tables)[0]
        result = result_rows_from_tables(tables)[0]
        self.assertEqual("EP-900", log["experiment_id"])
        self.assertEqual(70, log["temperature_C"])
        self.assertEqual("Feed started cleanly.", log["observation"])
        self.assertEqual("DLS particle size", result["measurement_type"])
        self.assertEqual(185, result["numeric_value"])
        self.assertEqual(185, result["value"])

    def test_plot_studio_formulas_follow_selected_run_and_metric(self) -> None:
        process = plot_studio_process_formula()
        measurement = plot_studio_measurement_formula()
        self.assertIn("'Bench Log'!$A$2:$A$1000=$B$3", process)
        self.assertIn("MATCH($B$4", process)
        self.assertIn("'Measurements'!$A$2:$A$1000=$B$3", measurement)
        self.assertIn("'Measurements'!$C$2:$C$1000=$B$5", measurement)

    def test_reaction_master_rolls_up_entry_tables_without_manual_copying(self) -> None:
        planned = reaction_master_excel_formula("Planned mass (g)", 2)
        activity = reaction_master_excel_formula("Last activity", 2)
        next_action = reaction_master_excel_formula("Next action", 2)
        google = google_reaction_master_array_formulas()
        self.assertIn("SUMIF('Batch Builder'!$A$2:$A$1000,$A2", planned)
        self.assertIn("'Bench Log'!$A$2:$A$1000", activity)
        self.assertIn("Measurements!$A$2:$A$1000", activity)
        self.assertIn("Enter batch quantities", next_action)
        self.assertIn("MAP(A2:A1000", google["Planned mass (g)"])
        self.assertIn("COUNTIF('Bench Log'!A2:A1000", google["Bench entries"])
        self.assertIn("Experiments!$A$2:$AB$1000,24", reaction_master_excel_formula("Mass plan status", 2))
        self.assertIn("'Batch Builder'!$AR$2:$AR$1000", reaction_master_excel_formula("Calculation issues", 2))
        self.assertIn("CALCULATION_ISSUES", google["Mass plan status"])
        self.assertIn("MASS_MISMATCH", google["Mass plan status"])
        self.assertIn("'Batch Builder'!$AV$2:$AV$1000", reaction_master_excel_formula("Planned numerator eq", 2))
        self.assertIn("RATIO_MISMATCH", google["Stoichiometry status"])
        self.assertIn("Record Signatures", reaction_master_excel_formula("Signature status", 2))
        self.assertIn("Experiment Templates", reaction_master_excel_formula("Governance status", 2))
        self.assertIn("Notebook Sections", reaction_master_excel_formula("Section readiness", 2))
        self.assertIn("RAW_FILE_LINKS_MISSING", google["Raw data readiness"])
        self.assertIn("Reaction Outcomes", google["Outcome readiness"])
        self.assertIn("$AU2", reaction_master_excel_formula("Governance status", 2))
        self.assertIn("READY_TO_ARCHIVE", google["Governance status"])
        self.assertTrue(all(value.count("(") == value.count(")") for value in google.values()))

    def test_generated_workbook_prioritizes_compact_entry_and_plotting(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = save_workbook(Path(tmpdir) / "scientist.xlsx")
            workbook = load_workbook(path, data_only=False)
            self.assertEqual("visible", workbook["Bench Log"].sheet_state)
            self.assertEqual("visible", workbook["Measurements"].sheet_state)
            self.assertEqual("visible", workbook["Plot Studio"].sheet_state)
            self.assertEqual("visible", workbook["Reaction Master"].sheet_state)
            self.assertEqual("hidden", workbook["Daily Log"].sheet_state)
            self.assertEqual("hidden", workbook["Results"].sheet_state)
            self.assertEqual("hidden", workbook["Plot Dashboard"].sheet_state)
            self.assertEqual("Run ID", workbook["Bench Log"]["A1"].value)
            self.assertEqual("Observation / action", workbook["Bench Log"]["K1"].value)
            self.assertEqual("PLOT STUDIO", workbook["Plot Studio"]["A1"].value)
            self.assertEqual(2, len(workbook["Plot Studio"]._charts))
            self.assertIn("SUMIF", workbook["Reaction Master"]["J2"].value)
            self.assertIn("COUNTIF", workbook["Reaction Master"]["N2"].value)
            self.assertIn("Continue run", workbook["Reaction Master"]["R2"].value)
            self.assertIn("Experiments!$W2", workbook["Reaction Master"]["U2"].value)
            self.assertIn("MASS_MISMATCH", workbook["Reaction Master"]["AA2"].value)
            metadata = {
                row[0].value: row[1].value
                for row in workbook["Workbook Metadata"].iter_rows(min_row=2)
                if row[0].value
            }
            self.assertEqual("0.12.0", metadata["contract_version"])
            self.assertEqual("Stoichiometry status", workbook["Reaction Master"]["AK1"].value)
            self.assertEqual("Recovered mass (g)", workbook["Bench Log"]["R1"].value)
            self.assertEqual("Governance status", workbook["Reaction Master"]["AT1"].value)
            self.assertEqual("Section readiness", workbook["Reaction Master"]["AU1"].value)
            self.assertEqual("Raw data readiness", workbook["Reaction Master"]["AV1"].value)
            self.assertEqual("Outcome readiness", workbook["Reaction Master"]["AW1"].value)
            self.assertTrue({"Experiment Templates", "Inventory Transactions", "Equipment Bookings", "Record Signatures"}.issubset(workbook.sheetnames))
            self.assertTrue({"Notebook Sections", "Reaction Outcomes"}.issubset(workbook.sheetnames))
            self.assertEqual("visible", workbook["Notebook Sections"].sheet_state)
            self.assertEqual("visible", workbook["Reaction Outcomes"].sheet_state)
            self.assertIn("Batch Builder", workbook["Reaction Outcomes"]["L2"].value)
            self.assertEqual("0.0%", workbook["Reaction Outcomes"]["F2"].number_format)

    def test_reaction_outcomes_calculate_yield_and_mass_closure(self) -> None:
        self.assertIn("$E2/$D2", excel_reaction_outcome_formula("isolated_yield_percent", 2))
        self.assertIn("'Batch Builder'!$BB$2:$BB$1000", excel_reaction_outcome_formula("actual_input_mass_g", 2))
        self.assertIn("OVER_ACCOUNTED", excel_reaction_outcome_formula("mass_balance_status", 2))
        google = google_reaction_outcome_array_formulas()
        self.assertIn("MAP(A2:A1000", google["actual_input_mass_g"])
        self.assertIn("MISSING_TOLERANCE", google["mass_balance_status"])
        self.assertTrue(all(value.count("(") == value.count(")") for value in google.values()))

        requests = reaction_outcomes_setup_requests(900, is_new=True)
        self.assertTrue(any("formulaValue" in str(request) for request in requests))
        self.assertTrue(any("EP-001-PRODUCT" in str(request) for request in requests))
        self.assertEqual("content", sheet_by_name("Notebook Sections").headers[7])

    def test_record_governance_requires_effective_template_signoff_and_witness(self) -> None:
        experiment = {"experiment_id": "EP-10", "status": "complete", "template_id": "TPL-1", "template_version": "2"}
        templates = [{"template_id": "TPL-1", "version": "2", "state": "effective"}]
        unsigned = summarize_record_governance(experiment, templates, [], [])
        self.assertEqual("SIGNOFF_REQUIRED", unsigned["governance_status"])
        signed = [{"experiment_id": "EP-10", "status": "signed", "signer": "Scientist", "signed_at": "2026-08-12T12:00:00"}]
        unwitnessed = summarize_record_governance(experiment, templates, signed, [])
        self.assertEqual("WITNESS_REQUIRED", unwitnessed["governance_status"])
        signed[0].update({"witnessed_by": "Reviewer", "witnessed_at": "2026-08-12T13:00:00"})
        ready = summarize_record_governance(experiment, templates, signed, [{"experiment_id": "EP-10"}])
        self.assertEqual("READY_TO_ARCHIVE", ready["governance_status"])
        self.assertEqual(1, ready["material_transactions"])

    def test_live_upgrade_adds_entry_sheets_seeds_history_and_adds_two_charts(self) -> None:
        metadata = {
            "properties": {"timeZone": "America/Chicago"},
            "sheets": [
                {"properties": {"sheetId": index + 10, "title": title, "gridProperties": {"rowCount": 1000, "columnCount": 40}}}
                for index, title in enumerate(
                    (
                        "Run Console",
                        "Master Reagents",
                        "Experiments",
                        "Batch Builder",
                        "Daily Log",
                        "Formulations",
                        "Results",
                        "Plot Dashboard",
                    )
                )
            ],
        }
        tables = {
            "Daily Log": [{"experiment_id": "EP-001", "timestamp": "2026-08-12T10:00:00", "process_stage": "feed", "observation": "Stable."}],
            "Results": [{"experiment_id": "EP-001", "sample_id": "S1", "measurement_type": "DLS particle size", "numeric_value": 200, "units": "nm"}],
        }
        requests = google_scientist_workspace_upgrade_requests(metadata, tables)
        added_titles = {
            request["addSheet"]["properties"]["title"]
            for request in requests
            if "addSheet" in request
        }
        self.assertTrue({"Reaction Master", "Bench Log", "Measurements", "Plot Studio"}.issubset(added_titles))
        self.assertEqual(2, sum("addChart" in request for request in requests))
        self.assertTrue(any("Stable." in str(request) for request in requests))
        self.assertTrue(any("DLS particle size" in str(request) for request in requests))
        self.assertTrue(any("Bench Log" in str(request) for request in requests))
        self.assertTrue(any("Plot Studio" in str(request) for request in requests))
        self.assertTrue(any("Batch Builder" in str(request) and "formulaValue" in str(request) for request in requests))


if __name__ == "__main__":
    unittest.main()
