from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from lab_notebook_agent.cli import main
from lab_notebook_agent.google_api import run_live_google_plot_refresh
from lab_notebook_agent.google_sheets import batch_update_requests_from_report
from lab_notebook_agent.plotting import (
    apply_plot_report_to_workbook,
    build_plot_report,
    google_chart_requests,
)
from lab_notebook_agent.schema import SHEETS
from lab_notebook_agent.templates import save_workbook


def plotting_tables() -> dict[str, list[dict[str, Any]]]:
    return {
        "Experiments": [
            {"experiment_id": "EP-101", "date": "2026-07-01"},
            {"experiment_id": "EP-102", "date": "2026-07-02"},
        ],
        "Project Notebook Records": [
            {
                "record_id": "PN-PROC-1",
                "experiment_id": "EP-101",
                "record_type": "observation",
                "section": "process log",
                "elapsed_end_min": 0,
                "temperature_C": 70,
                "oil_temperature_C": 75,
                "rpm": 250,
                "torque": 80,
                "cumulative_addition_g": 0,
                "source_timestamp": "10:00 AM",
                "source_range": "Run!B10:I10",
                "active": "true",
            },
            {
                "record_id": "PN-PROC-2",
                "experiment_id": "EP-101",
                "record_type": "observation",
                "section": "process log",
                "elapsed_end_min": 15,
                "temperature_C": 72,
                "oil_temperature_C": 77,
                "rpm": 255,
                "torque": 82,
                "cumulative_addition_g": 15.5,
                "source_timestamp": "10:15 AM",
                "source_range": "Run!B11:I11",
                "active": "true",
            },
            {
                "record_id": "PN-FEED-1",
                "experiment_id": "EP-101",
                "record_type": "feed_step",
                "elapsed_start_min": 0,
                "elapsed_end_min": 15,
                "feed_rate_mL_min": 0.6,
                "cumulative_percent": 25,
                "radical_flux_mol_min_L": 0.0004,
                "particle_size_nm": 124,
                "source_range": "Run!C20:F20",
                "active": "true",
            },
            {
                "record_id": "PN-FEED-2",
                "experiment_id": "EP-101",
                "record_type": "feed_step",
                "elapsed_start_min": 15,
                "elapsed_end_min": 30,
                "feed_rate_mL_min": 0.75,
                "cumulative_percent": 50,
                "radical_flux_mol_min_L": 0.00035,
                "particle_size_nm": 142,
                "source_range": "Run!C21:F21",
                "active": "true",
            },
            {
                "record_id": "PN-COMP-1",
                "experiment_id": "EP-101",
                "record_type": "component",
                "material_name": "Butyl acrylate",
                "mass_g": 100,
                "actual_mass_g": 99.5,
                "parts_per_hundred_monomer": 90,
                "source_range": "Run!C4:H4",
                "active": "true",
            },
            {
                "record_id": "PN-COMP-2",
                "experiment_id": "EP-101",
                "record_type": "component",
                "material_name": "Surfactant",
                "mass_g": 2,
                "actual_mass_g": 2.1,
                "parts_per_hundred_monomer": 1.8,
                "source_range": "Run!C5:H5",
                "active": "true",
            },
        ],
        "Daily Log": [
            {
                "experiment_id": "EP-101",
                "timestamp": "2026-07-01T10:00:00",
                "process_stage": "feed",
                "temperature_C": 70,
                "rpm": 250,
                "solids_percent": 20,
                "particle_size_nm": 120,
                "observation": "Feed started.",
            },
            {
                "experiment_id": "EP-101",
                "timestamp": "2026-07-01T10:30:00",
                "process_stage": "feed",
                "temperature_C": 72,
                "rpm": 255,
                "solids_percent": 35,
                "conversion_percent": 80,
                "particle_size_nm": 140,
                "observation": "Feed halfway.",
            },
        ],
        "Results": [
            {
                "experiment_id": "EP-101",
                "sample_id": "EP-101-L1",
                "measurement_type": "DLS particle size",
                "value": 142,
                "units": "nm",
                "quality_flag": "ok",
            },
            {
                "experiment_id": "EP-102",
                "sample_id": "EP-102-L1",
                "measurement_type": "DLS particle size",
                "value": 128,
                "units": "nm",
                "quality_flag": "ok",
            },
        ],
    }


class PlotReportTests(unittest.TestCase):
    def test_builds_plot_ready_records_for_process_feed_charge_logs_and_results(
        self,
    ) -> None:
        report = build_plot_report(
            plotting_tables(),
            refreshed_at="2026-07-24T12:00:00+00:00",
        )
        kinds = {
            definition["plot_kind"]
            for definition in report["plot_definitions"]
        }
        self.assertTrue(
            {
                "process_temperature",
                "process_rpm",
                "process_torque",
                "cumulative_initiator",
                "feed_rate",
                "cumulative_feed",
                "radical_flux",
                "estimated_particle_size",
                "charge_mass_accuracy",
                "composition_pphm",
                "daily_temperature",
                "daily_percent_outcomes",
                "result_trend",
            }.issubset(kinds)
        )
        self.assertGreater(report["summary"]["ready_plot_count"], 10)
        data = report["plot_data_rows"]
        for definition in report["plot_definitions"]:
            start = int(definition["data_start_row"])
            end = int(definition["data_end_row"])
            self.assertEqual("header", data[start - 2]["row_type"])
            self.assertEqual(definition["plot_id"], data[start - 2]["plot_id"])
            self.assertEqual(
                definition["source_point_count"],
                end - start,
            )
        temperature = next(
            definition
            for definition in report["plot_definitions"]
            if definition["plot_kind"] == "process_temperature"
        )
        first_point = data[int(temperature["data_start_row"]) - 1]
        self.assertEqual("PN-PROC-1", first_point["source_record_ids"])
        self.assertEqual("Run!B10:I10", first_point["source_ranges"])
        self.assertIsInstance(first_point["x_value"], float)
        self.assertIsInstance(first_point["series_1_value"], float)

    def test_plot_report_is_idempotent_when_only_refresh_timestamps_change(self) -> None:
        tables = plotting_tables()
        first = build_plot_report(
            tables,
            refreshed_at="2026-07-24T12:00:00+00:00",
        )
        tables["Plot Data"] = first["plot_data_rows"]
        tables["Plot Definitions"] = first["plot_definitions"]
        tables["Plot Dashboard"] = first["plot_dashboard_rows"]
        for sheet_name in ("Plot Data", "Plot Definitions", "Plot Dashboard"):
            spec = next(spec for spec in SHEETS if spec.name == sheet_name)
            tables[sheet_name] = [
                {
                    header: (
                        str(row.get(header, ""))
                        if isinstance(row.get(header, ""), (int, float))
                        else row.get(header, "")
                    )
                    for header in spec.headers
                }
                for row in tables[sheet_name]
            ]
        second = build_plot_report(
            tables,
            refreshed_at="2026-07-25T12:00:00+00:00",
        )
        self.assertTrue(second["unchanged"])
        self.assertNotIn("replace_plot_data", second)
        self.assertNotIn("replace_plot_definitions", second)
        self.assertNotIn("replace_plot_dashboard", second)

    def test_incremental_initiator_readings_are_accumulated_for_plotting(self) -> None:
        report = build_plot_report(
            {
                "Project Notebook Records": [
                    {
                        "record_id": "PN-I-1",
                        "experiment_id": "BMS-018",
                        "record_type": "observation",
                        "section": "process log",
                        "elapsed_end_min": 0,
                        "actual_value": 10,
                        "units": "g",
                        "active": "true",
                    },
                    {
                        "record_id": "PN-I-2",
                        "experiment_id": "BMS-018",
                        "record_type": "observation",
                        "section": "process log",
                        "elapsed_end_min": 5,
                        "actual_value": 20,
                        "units": "g",
                        "active": "true",
                    },
                ]
            }
        )
        definition = next(
            definition
            for definition in report["plot_definitions"]
            if definition["plot_kind"] == "cumulative_initiator"
        )
        start = int(definition["data_start_row"])
        values = [
            report["plot_data_rows"][row_number - 2]["series_1_value"]
            for row_number in range(start + 1, int(definition["data_end_row"]) + 1)
        ]
        self.assertEqual([10.0, 30.0], values)

    def test_local_workbook_contains_managed_data_definitions_and_charts(self) -> None:
        report = build_plot_report(plotting_tables())
        with tempfile.TemporaryDirectory() as tmpdir:
            workbook_path = save_workbook(
                Path(tmpdir) / "notebook.xlsx",
                include_examples=False,
            )
            apply_plot_report_to_workbook(workbook_path, report)
            workbook = load_workbook(workbook_path, data_only=False)
            self.assertEqual(
                len(report["plot_data_rows"]) + 1,
                workbook["Plot Data"].max_row,
            )
            self.assertEqual(
                len(report["plot_definitions"]) + 1,
                workbook["Plot Definitions"].max_row,
            )
            self.assertEqual(
                report["summary"]["ready_plot_count"],
                len(workbook["Plot Dashboard"]._charts),
            )

    def test_local_plot_cli_reads_canonical_tabs_and_applies_dashboard(self) -> None:
        tables = plotting_tables()
        with tempfile.TemporaryDirectory() as tmpdir:
            workbook_path = save_workbook(
                Path(tmpdir) / "notebook.xlsx",
                include_examples=False,
            )
            workbook = load_workbook(workbook_path)
            for sheet_name in (
                "Experiments",
                "Project Notebook Records",
                "Daily Log",
                "Results",
            ):
                worksheet = workbook[sheet_name]
                headers = [str(cell.value or "") for cell in worksheet[1]]
                for row in tables.get(sheet_name, []):
                    worksheet.append([row.get(header, "") for header in headers])
            workbook.save(workbook_path)
            report_path = Path(tmpdir) / "plot-report.json"
            self.assertEqual(
                0,
                main(
                    [
                        "plot-notebook",
                        "--workbook",
                        str(workbook_path),
                        "--apply",
                        "--report-output",
                        str(report_path),
                    ]
                ),
            )
            workbook = load_workbook(workbook_path, data_only=False)
            self.assertGreater(len(workbook["Plot Dashboard"]._charts), 10)
            self.assertTrue(report_path.exists())
            second_report_path = Path(tmpdir) / "plot-report-second.json"
            self.assertEqual(
                0,
                main(
                    [
                        "plot-notebook",
                        "--workbook",
                        str(workbook_path),
                        "--report-output",
                        str(second_report_path),
                    ]
                ),
            )
            second = json.loads(second_report_path.read_text(encoding="utf-8"))
            self.assertTrue(second["unchanged"])


class GooglePlotTests(unittest.TestCase):
    def test_plot_replacement_requests_write_numbers_as_numbers_and_clear_tail(
        self,
    ) -> None:
        report = build_plot_report(plotting_tables())
        report["existing_plot_row_counts"]["Plot Data"] = (
            len(report["plot_data_rows"]) + 3
        )
        sheet_ids = {
            "Plot Data": 901,
            "Plot Definitions": 902,
            "Plot Dashboard": 903,
        }
        requests = batch_update_requests_from_report(report, sheet_ids)
        plot_update = next(
            request["updateCells"]
            for request in requests
            if request.get("updateCells", {})
            .get("range", {})
            .get("sheetId")
            == 901
        )
        number_cells = [
            cell
            for row in plot_update["rows"]
            for cell in row["values"]
            if "numberValue" in cell.get("userEnteredValue", {})
        ]
        self.assertTrue(number_cells)
        self.assertEqual(
            report["existing_plot_row_counts"]["Plot Data"],
            plot_update["range"]["endRowIndex"] - 1,
        )
        self.assertEqual(
            {},
            plot_update["rows"][-1]["values"][0],
        )

    def test_google_charts_are_added_updated_and_stale_managed_charts_deleted(
        self,
    ) -> None:
        report = build_plot_report(plotting_tables())
        first = next(
            definition
            for definition in report["plot_definitions"]
            if definition["status"] == "ready"
        )
        metadata = {
            "sheets": [
                {
                    "properties": {"title": "Plot Dashboard", "sheetId": 903},
                    "charts": [
                        {
                            "chartId": 7001,
                            "spec": {
                                "title": f"Old title [LNA:{first['plot_id']}]"
                            },
                        },
                        {
                            "chartId": 7002,
                            "spec": {"title": "Stale [LNA:plot-stale]"},
                        },
                    ],
                }
            ]
        }
        requests = google_chart_requests(
            report,
            {"Plot Data": 901, "Plot Dashboard": 903},
            metadata,
        )
        self.assertTrue(
            any(
                request.get("updateChartSpec", {}).get("chartId") == 7001
                for request in requests
            )
        )
        self.assertTrue(
            any(
                request.get("updateEmbeddedObjectPosition", {}).get("objectId")
                == 7001
                for request in requests
            )
        )
        self.assertTrue(any("addChart" in request for request in requests))
        self.assertTrue(
            any(
                request.get("deleteEmbeddedObject", {}).get("objectId") == 7002
                for request in requests
            )
        )
        added_spec = next(
            request["addChart"]["chart"]["spec"]
            for request in requests
            if "addChart" in request
        )
        domain = added_spec["basicChart"]["domains"][0]["domain"][
            "sourceRange"
        ]["sources"][0]
        self.assertEqual(901, domain["sheetId"])
        self.assertGreater(domain["endRowIndex"], domain["startRowIndex"])
        result_definition = next(
            definition
            for definition in report["plot_definitions"]
            if definition["plot_kind"] == "result_trend"
        )
        result_spec = next(
            request["addChart"]["chart"]["spec"]
            for request in requests
            if request.get("addChart", {})
            .get("chart", {})
            .get("spec", {})
            .get("title", "")
            .endswith(f"[LNA:{result_definition['plot_id']}]")
        )
        result_domain = result_spec["basicChart"]["domains"][0]["domain"][
            "sourceRange"
        ]["sources"][0]
        self.assertEqual(7, result_domain["startColumnIndex"])
        for request in requests:
            spec = (
                request.get("addChart", {})
                .get("chart", {})
                .get("spec", {})
            )
            basic_chart = spec.get("basicChart", {})
            if basic_chart.get("chartType") == "LINE":
                self.assertNotIn("stackedType", basic_chart)
            if basic_chart.get("chartType") == "COLUMN":
                self.assertEqual("NOT_STACKED", basic_chart.get("stackedType"))

    def test_live_plot_refresh_audits_and_writes_only_the_target(self) -> None:
        client = PlotSheetsClient(plotting_tables())
        run = run_live_google_plot_refresh(
            "target-notebook",
            client,
            apply=True,
        )
        self.assertTrue(run["apply_audit"]["valid"], run["apply_audit"])
        self.assertTrue(run["applied"])
        self.assertGreater(run["chart_request_count"], 0)
        self.assertEqual(1, len(client.batch_calls))
        self.assertEqual("target-notebook", client.batch_calls[0][0])
        self.assertTrue(
            any(
                "addChart" in request
                for request in client.batch_calls[0][1]
            )
        )


class PlotSheetsClient:
    def __init__(self, tables: dict[str, list[dict[str, Any]]]) -> None:
        self.metadata = {
            "properties": {"title": "Project Notebook"},
            "sheets": [
                {
                    "properties": {
                        "title": spec.name,
                        "sheetId": index,
                        "gridProperties": {
                            "rowCount": 1000,
                            "columnCount": len(spec.headers),
                        },
                    }
                }
                for index, spec in enumerate(SHEETS, start=100)
            ],
        }
        self.values: dict[str, list[list[Any]]] = {}
        for spec in SHEETS:
            rows = tables.get(spec.name, [])
            self.values[spec.name] = [
                list(spec.headers),
                *[
                    [row.get(header, "") for header in spec.headers]
                    for row in rows
                ],
            ]
        self.batch_calls: list[tuple[str, list[dict[str, Any]]]] = []

    def get_metadata(self, spreadsheet_id: str) -> dict[str, Any]:
        if spreadsheet_id != "target-notebook":
            raise KeyError(spreadsheet_id)
        return self.metadata

    def get_values(
        self,
        spreadsheet_id: str,
        sheet_name: str,
        value_range: str,
        value_render_option: str = "FORMATTED_VALUE",
    ) -> list[list[Any]]:
        if spreadsheet_id != "target-notebook":
            raise KeyError(spreadsheet_id)
        return self.values[sheet_name]

    def batch_update(
        self,
        spreadsheet_id: str,
        requests: list[dict[str, Any]],
    ) -> dict[str, Any]:
        self.batch_calls.append((spreadsheet_id, requests))
        return {"replies": [{} for _ in requests]}


if __name__ == "__main__":
    unittest.main()
