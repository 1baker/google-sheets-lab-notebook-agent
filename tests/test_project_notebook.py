from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from lab_notebook_agent.agent import is_current_experiment_row
from lab_notebook_agent.cli import main
from lab_notebook_agent.google_api import run_live_google_project_notebook_sync
from lab_notebook_agent.google_sheets import (
    audit_report_against_snapshot,
    batch_update_requests_from_report,
    snapshot_from_tables,
)
from lab_notebook_agent.project_notebook import (
    apply_project_notebook_sync_report_to_workbook,
    build_project_notebook_sync_report,
    parse_project_notebook_sheet,
    parse_project_notebook_workbook,
)
from lab_notebook_agent.schema import SHEETS
from lab_notebook_agent.sheets import (
    build_experiment_entry_from_tables,
    load_workbook_tables,
)
from lab_notebook_agent.templates import save_workbook


SOURCE_SPREADSHEET_ID = "1zhdi6LkRNhQRitZjzd1mkN4kOrtrMKiQ2Wx4H_Hws0c"


def variable_feed_values() -> list[list[Any]]:
    return [
        ["Core Stage"],
        ["", "", "Material", "Parts per hundred monomer", "Mass", "Density", "Volume"],
        ["", "Pre-Reactor Contents", "Seed Latex", 11, 12.1, 1, 12.1],
        ["", "", "DI Water", 50, 55, 1, 55],
        ["", "Monomer Pre-Emulsion Contents", "Butyl Acrylate", 99.7, 109.67, 0.89, 123.22],
        ["", "", "Aerosol MA-80", 1.6, 1.76, 1, 1.76],
        ["Final temperature", 70],
        ["Stir speed RPM", 250],
        ["Seed temperature peaked at 72 C, then equilibrated at 70 C before the feed began."],
        [
            "",
            "Core Monomer Time Feeds",
            "Time Intervals",
            "Core Monomer Feed Rate (mL/min)",
            "Instantaneous Radical Flux (mol/min/L)",
            "Estimated Particle Size (nm)",
        ],
        ["", "", "0-15min", 0.601, 0.0004, 124],
        ["", "", "15-30min", 0.75, 0.00035, 142],
    ]


def early_ep_values() -> list[list[Any]]:
    return [
        ["Seed Stage"],
        ["Material", "Parts per hundred monomer", "Mass", "wt %", "Density", "Volume"],
        ["DI Water", 90, 90, 60, 1, 90],
        ["Sodium dodecyl sulfate", 1.2, 1.2, 0.8, 1.01, 1.19],
        ["Ammonium persulfate", 0.4, 0.4, 0.27, 1, 0.4],
        ["Core Stage"],
        ["Material", "Parts per hundred monomer", "Mass", "wt %", "Density", "Volume"],
        ["Butyl acrylate", 85, 85, 56.7, 0.89, 95.5],
    ]


def ledger_values() -> list[list[Any]]:
    return [
        ["BM-PEG40-01"],
        ["Date made", "2025-06-13"],
        ["component", "wt. %", "goal, g", "actual, g"],
        ["xanthan", 0.001, 4, 3.999],
        ["PEG40 castor oil", 0.015, 60, 60.26],
        ["BioMAG emulsion", 0.3, 1200, 1198.4],
        ["total", 1, 4000, 3998.2],
        ["High shear blend for five minutes before the second water addition."],
        ["Slight separation overnight; shaking redistributed the emulsion."],
    ]


def charge_sheet_values() -> list[list[Any]]:
    return [
        ["", "Scale and Composition", "", "", "", "Step 1 : Mass Charge Values", "", "Actual", "Mass left"],
        ["", "AESO", 11281, "g", "", "AESO_SESO", 22562.0, 22590, -28.0, "g"],
        ["", "MW Target", 20000, "g/mol", "", "AG", 26234.8, 26235, -0.2, "g"],
        ["", "Initiator Ratio", 0.35, "mol I / mol CTA", "", "PEG-400", 39392, 39410, -18, "g"],
        ["", "", "", "", "", "CTA", 623.48, 624.53, -1.05, "g"],
        ["", "", "", "", "", "Step 2 : Initiate Polymerization"],
        ["", "", "", "", "", "AMBN in DES", 805.73, 800.9, 4.83, "g"],
        [
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "Step 2: When temperature is reached, charge with initiator.\nContinue Argon flush throughout reaction.",
        ],
        ["", "Time", "Duration", "Rxn Temp (C)", "RPM", "Oil Temp (C)", "Torque (W)", "initiator soln added"],
        ["", "11:10:56 AM", "", 35.4, 0, 78.7, "", "Heating started"],
        ["", "2:38:04 PM", "0:00:00", 94.8, 350, 104.5, 80.54, 160.6],
        ["", "2:43:01 PM", "0:04:57", 94.8, 350, 104.7, 81.15, 161.6],
        ["", "3:47:29 PM", "1:09:25", "", "", "", "", "kill"],
    ]


def charge_index_values() -> list[list[Any]]:
    return [
        ["Experiment Name/Description", "Date", "Lot Number"],
        ["", "", "AG", "AEHOSO/SESO", "CTA", "AMBN", "Solvent 1", "Solvent 2", "Product", "Notes"],
        [
            "Surfactant Batch #18",
            "08/11/2023",
            "AG-300-01-052323",
            "AHS-300-07-071923",
            "DESOX-10-04-062123",
            "I-01-061623",
            "PEG-02-080823",
            "OA-01-061623",
            "BMS-18-081123",
            "",
        ],
        [
            "DES-OX Surf #12",
            "7/20/2023",
            "AG-300-01-052323",
            "AHS-50-04-063023",
            "DESOX-10-02-051223",
            "I-01-061623",
            "N/A",
            "N/A",
            "N/A",
            "Polysorbate based run that worked best. Run for 80 mins",
        ],
    ]


def parse_recipe(
    values: list[list[Any]],
    sheet_name: str = "EP-324 Variable Feed Rates SFS Trial 2",
    sheet_id: int = 324,
    synced_at: str = "2026-07-23T12:00:00+00:00",
) -> dict[str, Any]:
    return parse_project_notebook_sheet(
        values,
        source_spreadsheet_id=SOURCE_SPREADSHEET_ID,
        source_title="Emulsion Polymerization",
        source_sheet_name=sheet_name,
        source_sheet_id=sheet_id,
        source_range="A1:AO220",
        synced_at=synced_at,
    )


class ProjectNotebookParserTests(unittest.TestCase):
    def test_parses_variable_feed_recipe_components_metrics_notes_and_schedule(self) -> None:
        parsed = parse_recipe(variable_feed_values())
        self.assertEqual("EP-324", parsed["experiment"]["experiment_id"])
        self.assertEqual("cochran-emulsion", parsed["parser_profile"])
        records = parsed["records"]
        components = [row for row in records if row["record_type"] == "component"]
        self.assertEqual(4, len(components))
        butyl_acrylate = next(row for row in components if row["material_name"] == "Butyl Acrylate")
        self.assertEqual("core_monomer", butyl_acrylate["target_role"])
        self.assertEqual("109.67", butyl_acrylate["mass_g"])
        self.assertEqual("99.7", butyl_acrylate["parts_per_hundred_monomer"])

        feed_steps = [row for row in records if row["record_type"] == "feed_step"]
        self.assertEqual(2, len(feed_steps))
        self.assertEqual("0", feed_steps[0]["elapsed_start_min"])
        self.assertEqual("15", feed_steps[0]["elapsed_end_min"])
        self.assertEqual("0.601", feed_steps[0]["feed_rate_mL_min"])
        self.assertEqual("124", feed_steps[0]["particle_size_nm"])
        self.assertTrue(any(row["temperature_C"] == "70" for row in records))
        self.assertTrue(any("equilibrated" in row["notes"] for row in records))
        self.assertTrue(all(row["record_id"].startswith("PN-") for row in records))
        self.assertTrue(all(row["source_range"] for row in records))

    def test_parses_early_ep_200_vg_layout_and_roles(self) -> None:
        parsed = parse_recipe(early_ep_values(), "EP-200-VG", 200)
        self.assertEqual("EP-200-VG", parsed["experiment"]["experiment_id"])
        components = [row for row in parsed["records"] if row["record_type"] == "component"]
        self.assertEqual(4, len(components))
        roles = {row["material_name"]: row["target_role"] for row in components}
        self.assertEqual("surfactant", roles["Sodium dodecyl sulfate"])
        self.assertEqual("initiator", roles["Ammonium persulfate"])
        self.assertEqual("core_monomer", roles["Butyl acrylate"])

    def test_parses_biomag_ledger_planned_actual_and_observations(self) -> None:
        parsed = parse_project_notebook_sheet(
            ledger_values(),
            source_spreadsheet_id="biomag-workbook",
            source_title="Biomag emulsion Recipe.xlsx",
            source_sheet_name="PEG40 Batch",
            source_sheet_id=9,
            source_range="A1:H100",
            synced_at="2026-07-23T12:00:00+00:00",
        )
        self.assertEqual("cochran-emulsion-ledger", parsed["parser_profile"])
        components = [row for row in parsed["records"] if row["record_type"] == "component"]
        self.assertEqual(3, len(components))
        xanthan = next(row for row in components if row["material_name"] == "xanthan")
        self.assertEqual("4", xanthan["planned_value"])
        self.assertEqual("3.999", xanthan["actual_value"])
        self.assertEqual("4", xanthan["mass_g"])
        self.assertEqual("3.999", xanthan["actual_mass_g"])
        observations = [row for row in parsed["records"] if row["record_type"] == "observation"]
        self.assertTrue(any("High shear" in row["notes"] for row in observations))
        self.assertTrue(any("separation overnight" in row["notes"] for row in observations))

    def test_parses_biomag_charge_targets_actuals_instructions_and_process_log(self) -> None:
        parsed = parse_project_notebook_sheet(
            charge_sheet_values(),
            source_spreadsheet_id="biomag-charge",
            source_title="BioMAG Surfactant Charge Sheet",
            source_sheet_name="300L Surfactant Batch #18",
            source_sheet_id=18,
            source_range="A1:AE160",
            synced_at="2026-07-23T12:00:00+00:00",
        )
        self.assertEqual("cochran-charge-sheet", parsed["parser_profile"])
        self.assertEqual("BMS-018", parsed["experiment"]["experiment_id"])
        components = [row for row in parsed["records"] if row["record_type"] == "component"]
        self.assertEqual(5, len(components))
        cta = next(row for row in components if row["material_name"] == "CTA")
        self.assertEqual("623.48", cta["mass_g"])
        self.assertEqual("624.53", cta["actual_mass_g"])
        self.assertEqual("-1.05", cta["mass_variance_g"])
        ambn = next(row for row in components if row["material_name"] == "AMBN in DES")
        self.assertEqual("initiator", ambn["target_role"])

        logs = [
            row
            for row in parsed["records"]
            if row["section"] == "process log"
        ]
        self.assertEqual(4, len(logs))
        self.assertEqual("94.8", logs[1]["temperature_C"])
        self.assertEqual("350", logs[1]["rpm"])
        self.assertEqual("104.5", logs[1]["oil_temperature_C"])
        self.assertEqual("80.54", logs[1]["torque"])
        self.assertEqual("W", logs[1]["torque_units"])
        self.assertEqual("160.6", logs[1]["actual_value"])
        self.assertEqual("4.95", logs[2]["elapsed_end_min"])
        self.assertTrue(any("Heating started" in row["notes"] for row in logs))
        self.assertTrue(
            any(
                "Continue Argon flush" in row["notes"]
                for row in parsed["records"]
                if row["record_type"] == "observation"
            )
        )

    def test_parses_cumulative_initiator_addition_for_process_plots(self) -> None:
        values = charge_sheet_values()
        values[8].append("Total initiator added")
        values[9].append("")
        values[10].append(160.6)
        values[11].append(322.2)
        values[12].append(322.2)
        parsed = parse_project_notebook_sheet(
            values,
            source_spreadsheet_id="biomag-charge",
            source_title="BioMAG Surfactant Charge Sheet",
            source_sheet_name="300L Surfactant Batch #18",
            source_sheet_id=18,
            source_range="A1:AE160",
            synced_at="2026-07-23T12:00:00+00:00",
        )
        logs = [
            row
            for row in parsed["records"]
            if row["section"] == "process log"
        ]
        self.assertEqual("160.6", logs[1]["cumulative_addition_g"])
        self.assertEqual("322.2", logs[2]["cumulative_addition_g"])

    def test_parses_charge_master_material_and_product_lot_lineage(self) -> None:
        parsed = parse_project_notebook_sheet(
            charge_index_values(),
            source_spreadsheet_id="biomag-charge",
            source_title="BioMAG Surfactant Charge Sheet",
            source_sheet_name="Master Sheet",
            source_sheet_id=375394058,
            source_range="A1:Z100",
            synced_at="2026-07-23T12:00:00+00:00",
        )
        self.assertEqual("cochran-charge-index", parsed["parser_profile"])
        self.assertEqual(
            {"BMS-018", "DESOX-012"},
            {
                experiment["experiment_id"]
                for experiment in parsed["experiments"]
            },
        )
        ag_lot = next(
            row
            for row in parsed["records"]
            if row["experiment_id"] == "BMS-018"
            and row["material_name"] == "AG"
        )
        self.assertEqual("AG-300-01-052323", ag_lot["lot_number"])
        product = next(
            row
            for row in parsed["records"]
            if row["experiment_id"] == "BMS-018"
            and row["label"] == "Product lot"
        )
        self.assertEqual("BMS-18-081123", product["product_lot"])
        self.assertTrue(
            any(
                "worked best" in row["notes"]
                for row in parsed["records"]
                if row["record_type"] == "observation"
            )
        )
        report = build_project_notebook_sync_report(
            [parsed],
            {spec.name: [] for spec in SHEETS},
        )
        self.assertEqual(2, report["summary"]["experiment_rows_to_append"])


class ProjectNotebookSyncTests(unittest.TestCase):
    def test_imported_records_are_excluded_from_same_experiment_context_search(self) -> None:
        self.assertTrue(
            is_current_experiment_row(
                {
                    "sheet": "Project Notebook Records",
                    "row": {"experiment_id": "EP-324"},
                },
                "EP-324",
            )
        )

    def test_sync_is_idempotent_and_retires_removed_source_records(self) -> None:
        parsed = parse_recipe(variable_feed_values())
        empty_tables = {spec.name: [] for spec in SHEETS}
        first = build_project_notebook_sync_report([parsed], empty_tables)
        self.assertEqual(1, first["summary"]["experiment_rows_to_append"])
        self.assertGreater(first["summary"]["project_records_to_append"], 0)
        tables = dict(empty_tables)
        tables["Experiments"] = list(first["append_experiments"])
        tables["Project Notebook Records"] = list(first["append_project_notebook_records"])
        tables["Source Sync"] = list(first["append_source_sync"])

        same_source = parse_recipe(
            variable_feed_values(),
            synced_at="2026-07-23T13:00:00+00:00",
        )
        second = build_project_notebook_sync_report([same_source], tables)
        self.assertEqual(1, second["summary"]["unchanged_source_count"])
        self.assertEqual(0, second["summary"]["project_records_to_append"])
        self.assertEqual(0, second["summary"]["project_record_cells_to_update"])

        changed_values = variable_feed_values()
        changed_values[2][4] = 13.5
        changed_values.pop(3)
        changed = parse_recipe(
            changed_values,
            synced_at="2026-07-23T14:00:00+00:00",
        )
        third = build_project_notebook_sync_report([changed], tables)
        self.assertEqual(1, third["summary"]["changed_source_count"])
        self.assertTrue(
            any(
                update["field"] == "mass_g" and update["value"] == "13.5"
                for update in third["update_project_notebook_records"]
            )
        )
        self.assertTrue(
            any(
                update["field"] == "active" and update["value"] == "false"
                for update in third["update_project_notebook_records"]
            )
        )

    def test_sync_report_audits_and_generates_google_requests(self) -> None:
        parsed = parse_recipe(variable_feed_values())
        tables = {spec.name: [] for spec in SHEETS}
        report = build_project_notebook_sync_report([parsed], tables)
        sheet_ids = {spec.name: index for index, spec in enumerate(SHEETS, start=100)}
        snapshot = snapshot_from_tables(tables, sheet_ids)
        audit = audit_report_against_snapshot(report, snapshot, require_sheet_ids=True)
        self.assertTrue(audit["valid"], audit)
        requests = batch_update_requests_from_report(report, sheet_ids)
        appended_sheet_ids = {
            request["appendCells"]["sheetId"]
            for request in requests
            if "appendCells" in request
        }
        self.assertIn(sheet_ids["Experiments"], appended_sheet_ids)
        self.assertIn(sheet_ids["Project Notebook Records"], appended_sheet_ids)
        self.assertIn(sheet_ids["Source Sync"], appended_sheet_ids)

    def test_sync_refuses_ambiguous_duplicate_target_keys(self) -> None:
        parsed = parse_recipe(variable_feed_values())
        tables = {spec.name: [] for spec in SHEETS}
        tables["Source Sync"] = [
            {"source_key": "duplicate"},
            {"source_key": "duplicate"},
        ]
        with self.assertRaisesRegex(ValueError, "Duplicate existing source_key"):
            build_project_notebook_sync_report([parsed], tables)

    def test_local_workbook_apply_adds_records_and_entry_context(self) -> None:
        parsed = parse_recipe(variable_feed_values())
        with tempfile.TemporaryDirectory() as tmpdir:
            workbook_path = save_workbook(Path(tmpdir) / "notebook.xlsx", include_examples=False)
            tables = load_workbook_tables(workbook_path)
            report = build_project_notebook_sync_report([parsed], tables)
            output = Path(tmpdir) / "synced.xlsx"
            apply_project_notebook_sync_report_to_workbook(
                workbook_path,
                report,
                output,
            )
            synced_tables = load_workbook_tables(output)
            self.assertEqual(1, len(synced_tables["Source Sync"]))
            self.assertGreater(len(synced_tables["Project Notebook Records"]), 5)
            entry = build_experiment_entry_from_tables(synced_tables, "EP-324")
            self.assertGreater(len(entry["formulation"]), 3)
            self.assertEqual(
                len(synced_tables["Project Notebook Records"]),
                len(entry["process_records"]),
            )
            workbook = load_workbook(output, read_only=True)
            self.assertIn("Project Notebook Records", workbook.sheetnames)
            self.assertIn("Source Sync", workbook.sheetnames)

    def test_local_excel_source_has_cli_dry_run_and_apply_path(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            source_path = Path(tmpdir) / "Biomag emulsion Recipe.xlsx"
            source_workbook = Workbook()
            source_sheet = source_workbook.active
            source_sheet.title = "PEG40 Batch"
            for row in ledger_values():
                source_sheet.append(row)
            source_workbook.save(source_path)
            target_path = save_workbook(
                Path(tmpdir) / "notebook.xlsx",
                include_examples=False,
            )
            report_path = Path(tmpdir) / "sync.json"

            parsed = parse_project_notebook_workbook(
                source_path,
                ("PEG40 Batch",),
                synced_at="2026-07-23T12:00:00+00:00",
            )
            self.assertEqual("cochran-emulsion-ledger", parsed[0]["parser_profile"])
            self.assertEqual(3, len([row for row in parsed[0]["records"] if row["record_type"] == "component"]))

            exit_code = main(
                [
                    "project-notebook-sync",
                    "--source-workbook",
                    str(source_path),
                    "--target-workbook",
                    str(target_path),
                    "--source-sheet",
                    "PEG40 Batch",
                    "--apply",
                    "--report-output",
                    str(report_path),
                ]
            )
            self.assertEqual(0, exit_code)
            report = json.loads(report_path.read_text(encoding="utf-8"))
            self.assertTrue(report["applied"])
            synced_tables = load_workbook_tables(target_path)
            self.assertEqual(1, len(synced_tables["Source Sync"]))
            self.assertEqual(
                3,
                len(
                    [
                        row
                        for row in synced_tables["Project Notebook Records"]
                        if row["record_type"] == "component"
                    ]
                ),
            )
            second_report_path = Path(tmpdir) / "sync-second.json"
            self.assertEqual(
                0,
                main(
                    [
                        "project-notebook-sync",
                        "--source-workbook",
                        str(source_path),
                        "--target-workbook",
                        str(target_path),
                        "--source-sheet",
                        "PEG40 Batch",
                        "--report-output",
                        str(second_report_path),
                    ]
                ),
            )
            second_report = json.loads(
                second_report_path.read_text(encoding="utf-8")
            )
            self.assertEqual(1, second_report["summary"]["unchanged_source_count"])
            self.assertEqual(0, second_report["summary"]["project_records_to_append"])
            self.assertEqual(
                0,
                second_report["summary"]["project_record_cells_to_update"],
            )

    def test_local_cli_refuses_to_overwrite_its_source_workbook(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            source_path = Path(tmpdir) / "source.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "PEG40 Batch"
            for row in ledger_values():
                worksheet.append(row)
            workbook.save(source_path)
            with self.assertRaisesRegex(SystemExit, "source workbook is read-only"):
                main(
                    [
                        "project-notebook-sync",
                        "--source-workbook",
                        str(source_path),
                        "--target-workbook",
                        str(source_path),
                        "--source-sheet",
                        "PEG40 Batch",
                        "--apply",
                    ]
                )


class RoutingSheetsClient:
    def __init__(self) -> None:
        old_specs = [
            spec
            for spec in SHEETS
            if spec.name not in {"Project Notebook Records", "Source Sync"}
        ]
        self.source_metadata = {
            "properties": {"title": "Emulsion Polymerization"},
            "sheets": [
                {
                    "properties": {
                        "title": "EP-324 Variable Feed Rates SFS Trial 2",
                        "sheetId": 324,
                    }
                }
            ],
        }
        self.target_metadata = {
            "properties": {"title": "Project Notebook"},
            "sheets": [
                {
                    "properties": {
                        "title": spec.name,
                        "sheetId": index,
                        "gridProperties": {"rowCount": 1000, "columnCount": len(spec.headers)},
                    }
                }
                for index, spec in enumerate(old_specs, start=100)
            ],
        }
        self.target_values = {
            spec.name: [list(spec.headers)]
            for spec in old_specs
        }
        self.batch_calls: list[tuple[str, list[dict[str, Any]]]] = []
        self.value_reads: list[tuple[str, str, str]] = []

    def get_metadata(self, spreadsheet_id: str) -> dict[str, Any]:
        if spreadsheet_id == SOURCE_SPREADSHEET_ID:
            return self.source_metadata
        if spreadsheet_id == "target-notebook":
            return self.target_metadata
        raise KeyError(spreadsheet_id)

    def get_values(
        self,
        spreadsheet_id: str,
        sheet_name: str,
        value_range: str,
        value_render_option: str = "FORMATTED_VALUE",
    ) -> list[list[Any]]:
        self.value_reads.append((spreadsheet_id, sheet_name, value_range))
        if spreadsheet_id == SOURCE_SPREADSHEET_ID:
            return variable_feed_values()
        if spreadsheet_id == "target-notebook" and sheet_name in self.target_values:
            return self.target_values[sheet_name]
        raise KeyError((spreadsheet_id, sheet_name))

    def batch_update(
        self,
        spreadsheet_id: str,
        requests: list[dict[str, Any]],
    ) -> dict[str, Any]:
        self.batch_calls.append((spreadsheet_id, requests))
        return {"replies": [{} for _ in requests]}


class ProjectNotebookLiveTests(unittest.TestCase):
    def test_live_sync_reads_source_and_only_writes_target(self) -> None:
        client = RoutingSheetsClient()
        run = run_live_google_project_notebook_sync(
            SOURCE_SPREADSHEET_ID,
            "target-notebook",
            ("EP-324 Variable Feed Rates SFS Trial 2",),
            client,
            apply=True,
        )
        self.assertTrue(run["source_is_read_only"])
        self.assertTrue(run["applied"])
        self.assertTrue(run["apply_audit"]["valid"], run["apply_audit"])
        self.assertTrue(run["setup_required"])
        self.assertGreater(run["setup_request_count"], 0)
        self.assertEqual(1, len(client.batch_calls))
        self.assertEqual("target-notebook", client.batch_calls[0][0])
        self.assertFalse(
            any(call[0] == SOURCE_SPREADSHEET_ID for call in client.batch_calls)
        )
        self.assertEqual(
            1,
            sum(
                1
                for spreadsheet_id, sheet_name, _ in client.value_reads
                if spreadsheet_id == SOURCE_SPREADSHEET_ID
                and sheet_name == "EP-324 Variable Feed Rates SFS Trial 2"
            ),
        )
        self.assertGreater(run["sync_report"]["summary"]["project_records_to_append"], 5)
        self.assertGreater(run["plot_report"]["summary"]["plot_count"], 0)
        self.assertGreater(run["plot_data_request_count"], 0)
        self.assertGreater(run["chart_request_count"], 0)
        self.assertTrue(
            any(
                "addChart" in request
                for request in run["batch_update_requests"]
            )
        )

    def test_live_sync_requires_explicit_source_tabs(self) -> None:
        with self.assertRaisesRegex(ValueError, "At least one source sheet"):
            run_live_google_project_notebook_sync(
                SOURCE_SPREADSHEET_ID,
                "target-notebook",
                (),
                RoutingSheetsClient(),
            )

    def test_live_sync_refuses_to_write_into_source_spreadsheet(self) -> None:
        with self.assertRaisesRegex(ValueError, "must differ"):
            run_live_google_project_notebook_sync(
                SOURCE_SPREADSHEET_ID,
                SOURCE_SPREADSHEET_ID,
                ("EP-324 Variable Feed Rates SFS Trial 2",),
                RoutingSheetsClient(),
                apply=True,
            )


if __name__ == "__main__":
    unittest.main()
