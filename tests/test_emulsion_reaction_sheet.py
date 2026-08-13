from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from lab_notebook_agent.emulsion_reaction_sheet import (
    SOURCE_SHEET,
    build_ccsp_audit_report,
    calculate_reaction_plan,
    ccsp_52_charges,
    ccsp_source_audit,
    save_ccsp_reaction_workbook,
)


class EmulsionReactionSheetTests(unittest.TestCase):
    def test_ccsp_52_plan_reconciles_stage_totals_and_has_no_calculation_failures(self) -> None:
        plan = calculate_reaction_plan()

        self.assertTrue(plan["checks"]["ready"], plan["checks"])
        self.assertEqual([], plan["checks"]["failures"])
        self.assertAlmostEqual(196.56, plan["stages"]["Seed"]["total_mass_g"], places=6)
        self.assertAlmostEqual(144.618, plan["stages"]["Core"]["total_mass_g"], places=6)
        self.assertAlmostEqual(162.3588, plan["stages"]["Shell"]["total_mass_g"], places=6)
        self.assertAlmostEqual(201.9396, plan["stages"]["Functional Shell"]["total_mass_g"], places=6)

    def test_corrected_core_molar_ratios_use_nonzero_main_aps_feed(self) -> None:
        ratios = calculate_reaction_plan()["ratios"]

        self.assertAlmostEqual(0.0006836707862, ratios["core_main_initiator_moles"], places=12)
        self.assertAlmostEqual(684.5074752, ratios["core_monomer_to_initiator"], places=6)
        self.assertAlmostEqual(0.1300813, ratios["core_cta_to_initiator"], places=6)
        self.assertAlmostEqual(14.0992, ratios["core_crosslinker_to_cta"], places=4)

    def test_zero_mass_charge_cannot_inherit_another_materials_volume(self) -> None:
        charges = ccsp_52_charges()
        zero_charge = charges[0].__class__(
            stage="Seed",
            addition="Pre-reactor",
            material="Sodium Acetate Buffer",
            role="buffer",
            source_cell="D6",
            phr=0.0,
            stage_factor_g_per_phr=0.6,
            density_g_ml=1.0,
            active_fraction=1.0,
            nonvolatile=True,
        )
        plan = calculate_reaction_plan((zero_charge,))
        self.assertEqual(0.0, plan["charges"][0]["mass_g"])
        self.assertEqual(0.0, plan["charges"][0]["volume_ml"])

    def test_design_split_is_reported_as_actual_not_assumed_exact(self) -> None:
        plan = calculate_reaction_plan()
        self.assertAlmostEqual(1.0, sum(plan["actual_stage_fractions"].values()), places=12)
        self.assertNotEqual(plan["target_stage_fractions"], plan["actual_stage_fractions"])
        self.assertLess(plan["checks"]["maximum_stage_fraction_deviation"], 0.005)

    def test_feed_schedule_finishes_every_stream_without_overallocation(self) -> None:
        from lab_notebook_agent.emulsion_reaction_sheet import _feed_schedule_rows

        rows = _feed_schedule_rows()
        for stage in ("Core", "Shell", "Functional Shell"):
            stage_rows = [row for row in rows if row[0] == stage]
            self.assertEqual(1.0, stage_rows[-1][2])
            self.assertAlmostEqual(1.0, sum(row[3] for row in stage_rows), places=12)
            self.assertAlmostEqual(1.0, sum(row[4] for row in stage_rows), places=12)
            self.assertTrue(all(later[1] > earlier[1] for earlier, later in zip(stage_rows, stage_rows[1:])))

    def test_audit_names_every_confirmed_source_issue(self) -> None:
        findings = ccsp_source_audit()
        cells = {finding["source_cell"] for finding in findings}
        self.assertIn("CCSP-52!G6", cells)
        self.assertIn("CCSP-52!K35:K36", cells)
        self.assertIn("CCSP-52!F74", cells)
        self.assertTrue(any("solids" in row["finding"].lower() for row in findings))

        report = build_ccsp_audit_report()
        self.assertEqual(SOURCE_SHEET, report["source"]["sheet"])
        self.assertTrue(report["deterministic_result"]["checks"]["ready"])
        json.dumps(report)

    def test_generated_workbook_is_simple_color_coded_and_has_feed_rates(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = save_ccsp_reaction_workbook(Path(tmpdir) / "ccsp.xlsx")
            workbook = load_workbook(path, data_only=False)
            self.assertEqual(
                ["Reaction Plan", "Feed Schedule", "Checks", "Assumptions"],
                workbook.sheetnames,
            )

            recipe = workbook["Reaction Plan"]
            self.assertEqual("D6", recipe.freeze_panes)
            self.assertEqual("Stage", recipe["A5"].value)
            self.assertIn("IF", recipe["I6"].value)
            self.assertIn("IFERROR", recipe["J6"].value)
            self.assertEqual("00FFF2CC", recipe["D6"].fill.fgColor.rgb)
            self.assertEqual("00E2F0D9", recipe["I6"].fill.fgColor.rgb)

            feeds = workbook["Feed Schedule"]
            self.assertEqual("End time (min)", feeds["B5"].value)
            self.assertEqual("Emulsion rate (mL/min)", feeds["H5"].value)
            self.assertEqual("Oxidant rate (mL/min)", feeds["I5"].value)
            self.assertEqual("Reductant rate (mL/min)", feeds["J5"].value)
            self.assertIn("SUMIFS", feeds["G6"].value)
            self.assertIn("IFERROR", feeds["I6"].value)
            self.assertEqual("00FFF2CC", feeds["B6"].fill.fgColor.rgb)
            self.assertEqual("00E2F0D9", feeds["H6"].fill.fgColor.rgb)

            checks = workbook["Checks"]
            self.assertIn("SUMIF", checks["B6"].value)
            self.assertIn("ABS", checks["B12"].value)

            assumptions = workbook["Assumptions"]
            self.assertEqual("Status", assumptions["A5"].value)
            self.assertEqual("VERIFY", assumptions["A6"].value)


if __name__ == "__main__":
    unittest.main()
