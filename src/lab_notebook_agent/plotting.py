from __future__ import annotations

import hashlib
import math
import re
from datetime import date, datetime, time, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill

from .project_notebook import ensure_workbook_contract
from .schema import sheet_by_name


PLOT_DATA_SHEET = "Plot Data"
PLOT_DEFINITIONS_SHEET = "Plot Definitions"
PLOT_DASHBOARD_SHEET = "Plot Dashboard"
MANAGED_CHART_TITLE_RE = re.compile(r"\s*\[LNA:([^\]]+)\]\s*$")
PLOT_REPLACEMENT_KEYS = {
    PLOT_DATA_SHEET: "replace_plot_data",
    PLOT_DEFINITIONS_SHEET: "replace_plot_definitions",
    PLOT_DASHBOARD_SHEET: "replace_plot_dashboard",
}
SERIES_VALUE_COLUMNS = (
    "series_1_value",
    "series_2_value",
    "series_3_value",
    "series_4_value",
)
def utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def numeric_value(value: Any) -> float | None:
    if value in ("", None):
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        parsed = float(value)
        return parsed if math.isfinite(parsed) else None
    text = str(value).strip().replace(",", "")
    if not text:
        return None
    if text.endswith("%"):
        text = text[:-1].strip()
    try:
        parsed = float(text)
        return parsed if math.isfinite(parsed) else None
    except ValueError:
        match = re.fullmatch(r"[-+]?(?:\d+(?:\.\d*)?|\.\d+)", text)
        if not match:
            return None
        parsed = float(match.group(0))
        return parsed if math.isfinite(parsed) else None


def stable_id(prefix: str, *parts: Any) -> str:
    payload = "\x1f".join(str(part).strip().lower() for part in parts)
    return f"{prefix}-{hashlib.sha256(payload.encode('utf-8')).hexdigest()[:16]}"


def active_row(row: dict[str, Any]) -> bool:
    return str(row.get("active", "true")).strip().lower() not in {
        "false",
        "0",
        "no",
        "inactive",
    }


def tables_after_report(
    tables: dict[str, list[dict[str, Any]]],
    report: dict[str, Any],
) -> dict[str, list[dict[str, Any]]]:
    """Project an append/update report in memory without mutating caller data."""

    projected = {name: [dict(row) for row in rows] for name, rows in tables.items()}
    for sheet_name, report_key in (
        ("Experiments", "append_experiments"),
        ("Project Notebook Records", "append_project_notebook_records"),
        ("Source Sync", "append_source_sync"),
    ):
        projected.setdefault(sheet_name, []).extend(
            dict(row)
            for row in report.get(report_key, []) or []
            if isinstance(row, dict)
        )
    for sheet_name, report_key in (
        ("Experiments", "update_experiments"),
        ("Project Notebook Records", "update_project_notebook_records"),
        ("Source Sync", "update_source_sync"),
    ):
        rows = projected.setdefault(sheet_name, [])
        for update in report.get(report_key, []) or []:
            if not isinstance(update, dict):
                continue
            row_number = int(update.get("row_number", 0) or 0)
            field = str(update.get("field", ""))
            if row_number >= 2 and field and row_number - 2 < len(rows):
                rows[row_number - 2][field] = update.get("value", "")
    for sheet_name, report_key in PLOT_REPLACEMENT_KEYS.items():
        if report_key in report:
            projected[sheet_name] = [
                dict(row)
                for row in report.get(report_key, []) or []
                if isinstance(row, dict)
            ]
    return projected


def build_plot_report(
    tables: dict[str, list[dict[str, Any]]],
    refreshed_at: str | None = None,
) -> dict[str, Any]:
    refreshed_at = refreshed_at or utc_now()
    blocks: list[dict[str, Any]] = []
    blocks.extend(project_notebook_plot_blocks(tables))
    blocks.extend(daily_log_plot_blocks(tables))
    blocks.extend(result_trend_plot_blocks(tables))
    blocks.sort(
        key=lambda block: (
            str(block.get("experiment_id", "")),
            str(block.get("plot_kind", "")),
            str(block.get("title", "")),
        )
    )
    plot_ids = [str(block.get("plot_id", "")) for block in blocks]
    if len(plot_ids) != len(set(plot_ids)):
        raise ValueError("Plot builders produced duplicate stable plot IDs.")

    plot_data: list[dict[str, Any]] = []
    definitions: list[dict[str, Any]] = []
    for block in blocks:
        points = block.get("points", [])
        if not points:
            continue
        plot_id = str(block["plot_id"])
        data_start_row = len(plot_data) + 2
        header = blank_plot_data_row()
        header.update(
            {
                "row_type": "header",
                "plot_id": plot_id,
                "experiment_id": block.get("experiment_id", ""),
                "plot_kind": block.get("plot_kind", ""),
                "x_value": block.get("x_axis_title", ""),
                "x_label": block.get("x_axis_title", ""),
                "x_units": block.get("x_units", ""),
                "recorded_at": refreshed_at,
                "active": "true",
            }
        )
        for series_index, series in enumerate(block.get("series", [])[:4], start=1):
            header[f"series_{series_index}_name"] = series["name"]
            header[f"series_{series_index}_value"] = series["name"]
            header[f"series_{series_index}_units"] = series.get("units", "")
        plot_data.append(header)
        for point_index, point in enumerate(points, start=1):
            row = blank_plot_data_row()
            row.update(
                {
                    "row_type": "point",
                    "plot_id": plot_id,
                    "plot_point_id": point.get("plot_point_id")
                    or stable_id(
                        "pt",
                        plot_id,
                        point.get("source_record_ids", ""),
                        point.get("x_value", ""),
                        point_index,
                    ),
                    "experiment_id": point.get(
                        "experiment_id", block.get("experiment_id", "")
                    ),
                    "plot_kind": block.get("plot_kind", ""),
                    "x_order": point.get("x_order", point_index),
                    "x_value": point.get("x_value", point_index),
                    "x_label": point.get("x_label", point.get("x_value", point_index)),
                    "x_units": block.get("x_units", ""),
                    "stage": point.get("stage", ""),
                    "section": point.get("section", ""),
                    "source_record_ids": point.get("source_record_ids", ""),
                    "source_ranges": point.get("source_ranges", ""),
                    "quality_flag": point.get("quality_flag", "observed"),
                    "recorded_at": point.get("recorded_at", refreshed_at),
                    "active": "true",
                }
            )
            for series_index, series in enumerate(block.get("series", [])[:4], start=1):
                row[f"series_{series_index}_name"] = series["name"]
                row[f"series_{series_index}_value"] = point.get(
                    series["field"], ""
                )
                row[f"series_{series_index}_units"] = series.get("units", "")
            plot_data.append(row)

        chart_type = str(block.get("chart_type", "LINE")).upper()
        point_count = len(points)
        minimum_points = 1 if chart_type == "COLUMN" else 2
        status = "ready" if point_count >= minimum_points else "insufficient_data"
        definition = {
            "plot_id": plot_id,
            "experiment_id": block.get("experiment_id", ""),
            "plot_kind": block.get("plot_kind", ""),
            "title": block.get("title", ""),
            "chart_type": chart_type,
            "x_axis_title": block.get("x_axis_title", ""),
            "y_axis_title": block.get("y_axis_title", ""),
            "data_start_row": data_start_row,
            "data_end_row": len(plot_data) + 1,
            "source_point_count": point_count,
            "status": status,
            "updated_at": refreshed_at,
            "notes": block.get("notes", ""),
            "x_axis_mode": block.get(
                "x_axis_mode",
                "category" if chart_type == "COLUMN" else "numeric",
            ),
        }
        for series_index, series in enumerate(block.get("series", [])[:4], start=1):
            definition[f"series_{series_index}_name"] = series["name"]
            definition[f"series_{series_index}_units"] = series.get("units", "")
        definitions.append(definition)

    dashboard_rows = [
        {
            "plot_id": definition["plot_id"],
            "title": definition["title"],
            "experiment_id": definition["experiment_id"],
            "plot_kind": definition["plot_kind"],
            "chart_type": definition["chart_type"],
            "point_count": definition["source_point_count"],
            "updated_at": refreshed_at,
        }
        for definition in definitions
    ]
    existing_counts = {
        sheet_name: len(tables.get(sheet_name, []))
        for sheet_name in PLOT_REPLACEMENT_KEYS
    }
    desired_by_sheet = {
        PLOT_DATA_SHEET: plot_data,
        PLOT_DEFINITIONS_SHEET: definitions,
        PLOT_DASHBOARD_SHEET: dashboard_rows,
    }
    replace: dict[str, list[dict[str, Any]]] = {}
    changed_sheets: list[str] = []
    for sheet_name, report_key in PLOT_REPLACEMENT_KEYS.items():
        desired = desired_by_sheet[sheet_name]
        existing = tables.get(sheet_name, [])
        ignored = (
            {"recorded_at"}
            if sheet_name == PLOT_DATA_SHEET
            else {"updated_at"}
        )
        numeric_fields = {
            PLOT_DATA_SHEET: {
                "x_order",
                "x_value",
                *SERIES_VALUE_COLUMNS,
            },
            PLOT_DEFINITIONS_SHEET: {
                "data_start_row",
                "data_end_row",
                "source_point_count",
            },
            PLOT_DASHBOARD_SHEET: {"point_count"},
        }[sheet_name]
        fields = tuple(
            field
            for field in sheet_by_name(sheet_name).headers
            if field not in ignored
        )
        if semantic_rows(existing, fields, numeric_fields) != semantic_rows(
            desired, fields, numeric_fields
        ):
            replace[report_key] = desired
            changed_sheets.append(sheet_name)

    return {
        "schema": "lab-notebook-agent-plot-report.v1",
        "plot_data_rows": plot_data,
        "plot_definitions": definitions,
        "plot_dashboard_rows": dashboard_rows,
        **replace,
        "existing_plot_row_counts": existing_counts,
        "changed_sheets": changed_sheets,
        "unchanged": not changed_sheets,
        "summary": {
            "plot_count": len(definitions),
            "ready_plot_count": sum(
                1 for definition in definitions if definition["status"] == "ready"
            ),
            "insufficient_data_plot_count": sum(
                1
                for definition in definitions
                if definition["status"] == "insufficient_data"
            ),
            "plot_point_count": sum(
                int(definition["source_point_count"]) for definition in definitions
            ),
            "changed_sheet_count": len(changed_sheets),
        },
    }


def blank_plot_data_row() -> dict[str, Any]:
    return {header: "" for header in sheet_by_name(PLOT_DATA_SHEET).headers}


def semantic_rows(
    rows: Iterable[dict[str, Any]],
    fields: tuple[str, ...],
    numeric_fields: set[str],
) -> list[dict[str, Any]]:
    return [
        {
            key: normalize_compare_value(row.get(key, ""), key in numeric_fields)
            for key in fields
        }
        for row in rows
    ]


def normalize_compare_value(value: Any, numeric_expected: bool) -> Any:
    if numeric_expected:
        numeric = numeric_value(value)
        if numeric is not None:
            return numeric
    return "" if value is None else str(value).strip()


def project_notebook_plot_blocks(
    tables: dict[str, list[dict[str, Any]]],
) -> list[dict[str, Any]]:
    records = [
        row
        for row in tables.get("Project Notebook Records", [])
        if isinstance(row, dict) and active_row(row)
    ]
    by_experiment = group_rows(records, "experiment_id")
    blocks: list[dict[str, Any]] = []
    for experiment_id, experiment_rows in by_experiment.items():
        process_rows = [
            row
            for row in experiment_rows
            if str(row.get("record_type", "")).strip()
            in {"observation", "process_parameter"}
            and (
                any(
                    numeric_value(row.get(field)) is not None
                    for field in (
                        "temperature_C",
                        "rpm",
                        "oil_temperature_C",
                        "torque",
                        "cumulative_addition_g",
                    )
                )
                or (
                    "process log"
                    in str(row.get("section", "")).strip().lower()
                    and str(row.get("units", "")).strip().lower()
                    in {"g", "gram", "grams"}
                    and numeric_value(row.get("actual_value")) is not None
                )
            )
        ]
        process_points = process_plot_points(process_rows)
        torque_units = {
            str(point.get("torque_units", "")).strip()
            for point in process_points
            if str(point.get("torque_units", "")).strip()
        }
        torque_unit = (
            next(iter(torque_units))
            if len(torque_units) == 1
            else "source units"
        )
        blocks.extend(
            metric_blocks(
                experiment_id,
                process_points,
                (
                    (
                        "process_temperature",
                        "Reaction and oil temperature",
                        (
                            ("reaction_temperature", "Reaction temperature", "°C"),
                            ("oil_temperature", "Oil temperature", "°C"),
                        ),
                        "Temperature (°C)",
                    ),
                    (
                        "process_rpm",
                        "Agitation speed",
                        (("rpm_value", "RPM", "rpm"),),
                        "Agitation speed (rpm)",
                    ),
                    (
                        "process_torque",
                        "Agitator torque / power",
                        (("torque_value", "Torque / power", torque_unit),),
                        f"Torque / power ({torque_unit})",
                    ),
                    (
                        "cumulative_initiator",
                        "Cumulative initiator addition",
                        (
                            (
                                "cumulative_addition",
                                "Cumulative initiator added",
                                "g",
                            ),
                        ),
                        "Cumulative addition (g)",
                    ),
                ),
                x_axis_title="Elapsed process time (min)",
                x_units="min",
            )
        )

        feed_rows = [
            row
            for row in experiment_rows
            if str(row.get("record_type", "")).strip() == "feed_step"
            or any(
                numeric_value(row.get(field)) is not None
                for field in (
                    "feed_rate_mL_min",
                    "cumulative_percent",
                    "radical_flux_mol_min_L",
                    "particle_size_nm",
                )
            )
        ]
        feed_points = feed_plot_points(feed_rows)
        blocks.extend(
            metric_blocks(
                experiment_id,
                feed_points,
                (
                    (
                        "feed_rate",
                        "Feed rate profile",
                        (("feed_rate", "Feed rate", "mL/min"),),
                        "Feed rate (mL/min)",
                    ),
                    (
                        "cumulative_feed",
                        "Cumulative feed profile",
                        (("cumulative_percent", "Cumulative feed", "%"),),
                        "Cumulative feed (%)",
                    ),
                    (
                        "radical_flux",
                        "Radical flux profile",
                        (
                            (
                                "radical_flux",
                                "Radical flux",
                                "mol/min/L",
                            ),
                        ),
                        "Radical flux (mol/min/L)",
                    ),
                    (
                        "estimated_particle_size",
                        "Estimated particle-size profile",
                        (("particle_size", "Particle size", "nm"),),
                        "Particle size (nm)",
                    ),
                ),
                x_axis_title="Feed time (min)",
                x_units="min",
            )
        )

        component_rows = [
            row
            for row in experiment_rows
            if str(row.get("record_type", "")).strip() == "component"
        ]
        charge_points = component_plot_points(component_rows)
        charge_series = (
            {"field": "planned_mass", "name": "Planned mass", "units": "g"},
            {"field": "actual_mass", "name": "Actual mass", "units": "g"},
        )
        charge_block = make_block(
            experiment_id,
            "charge_mass_accuracy",
            f"{experiment_id} — planned vs actual charge mass",
            "COLUMN",
            "Material",
            "Mass (g)",
            "",
            charge_series,
            charge_points,
            "Actual and planned source charge masses remain separately auditable.",
        )
        if charge_block:
            blocks.append(charge_block)
        pphm_points = [
            point for point in charge_points if point.get("pphm_value") is not None
        ]
        pphm_block = make_block(
            experiment_id,
            "composition_pphm",
            f"{experiment_id} — formulation composition (pphm)",
            "COLUMN",
            "Material",
            "Parts per hundred monomer",
            "",
            ({"field": "pphm_value", "name": "Composition", "units": "pphm"},),
            pphm_points,
            "Source pphm values by material.",
        )
        if pphm_block:
            blocks.append(pphm_block)
    return blocks


def process_plot_points(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    if not rows:
        return []
    clock_minutes = [clock_minute(row.get("source_timestamp")) for row in rows]
    first_clock = next((value for value in clock_minutes if value is not None), None)
    points: list[dict[str, Any]] = []
    previous_elapsed = 0.0
    running_addition = 0.0
    has_running_addition = False
    for index, row in enumerate(rows, start=1):
        elapsed = numeric_value(row.get("elapsed_end_min"))
        if elapsed is None and clock_minutes[index - 1] is not None and first_clock is not None:
            elapsed = clock_minutes[index - 1] - first_clock
            if elapsed < 0:
                elapsed += 24 * 60
        if elapsed is None:
            elapsed = previous_elapsed if points else 0.0
        previous_elapsed = elapsed
        cumulative_addition = numeric_value(row.get("cumulative_addition_g"))
        if cumulative_addition is not None:
            running_addition = cumulative_addition
            has_running_addition = True
        else:
            units = str(row.get("units", "")).strip().lower()
            incremental_addition = (
                numeric_value(row.get("actual_value"))
                if units in {"g", "gram", "grams"}
                and "process log"
                in str(row.get("section", "")).strip().lower()
                else None
            )
            if incremental_addition is not None:
                running_addition += incremental_addition
                has_running_addition = True
                cumulative_addition = running_addition
            elif has_running_addition:
                cumulative_addition = running_addition
        point = source_point(row, index, elapsed)
        point.update(
            {
                "reaction_temperature": numeric_value(row.get("temperature_C")),
                "oil_temperature": numeric_value(row.get("oil_temperature_C")),
                "rpm_value": numeric_value(row.get("rpm")),
                "torque_value": numeric_value(row.get("torque")),
                "torque_units": row.get("torque_units", ""),
                "cumulative_addition": (
                    cumulative_addition if has_running_addition else None
                ),
            }
        )
        points.append(point)
    return points


def feed_plot_points(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    points: list[dict[str, Any]] = []
    for index, row in enumerate(rows, start=1):
        elapsed = numeric_value(row.get("elapsed_end_min"))
        if elapsed is None:
            elapsed = numeric_value(row.get("elapsed_start_min"))
        if elapsed is None:
            elapsed = float(index)
        point = source_point(row, index, elapsed)
        point.update(
            {
                "feed_rate": numeric_value(row.get("feed_rate_mL_min")),
                "cumulative_percent": numeric_value(row.get("cumulative_percent")),
                "radical_flux": numeric_value(row.get("radical_flux_mol_min_L")),
                "particle_size": numeric_value(row.get("particle_size_nm")),
            }
        )
        points.append(point)
    return points


def component_plot_points(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    points: list[dict[str, Any]] = []
    for index, row in enumerate(rows, start=1):
        planned = numeric_value(row.get("mass_g"))
        actual = numeric_value(row.get("actual_mass_g"))
        pphm = numeric_value(row.get("parts_per_hundred_monomer"))
        if planned is None and actual is None and pphm is None:
            continue
        point = source_point(row, len(points) + 1, float(len(points) + 1))
        point.update(
            {
                "x_label": row.get("material_name")
                or row.get("label")
                or f"Material {len(points) + 1}",
                "planned_mass": planned,
                "actual_mass": actual,
                "pphm_value": pphm,
                "quality_flag": "observed" if actual is not None else "planned",
            }
        )
        points.append(point)
    return points


def source_point(
    row: dict[str, Any],
    order: int,
    x_value: float,
) -> dict[str, Any]:
    return {
        "experiment_id": row.get("experiment_id", ""),
        "x_order": order,
        "x_value": x_value,
        "x_label": row.get("source_timestamp")
        or row.get("label")
        or str(x_value),
        "stage": row.get("stage", ""),
        "section": row.get("section", ""),
        "source_record_ids": row.get("record_id", ""),
        "source_ranges": row.get("source_range", ""),
        "quality_flag": "observed",
        "recorded_at": row.get("source_timestamp")
        or row.get("synced_at")
        or "",
    }


def metric_blocks(
    experiment_id: str,
    points: list[dict[str, Any]],
    metrics: tuple[
        tuple[str, str, tuple[tuple[str, str, str], ...], str],
        ...,
    ],
    *,
    x_axis_title: str,
    x_units: str,
) -> list[dict[str, Any]]:
    blocks: list[dict[str, Any]] = []
    for plot_kind, title, raw_series, y_axis_title in metrics:
        series = tuple(
            {"field": field, "name": name, "units": units}
            for field, name, units in raw_series
        )
        relevant = [
            point
            for point in points
            if any(point.get(item["field"]) is not None for item in series)
        ]
        block = make_block(
            experiment_id,
            plot_kind,
            f"{experiment_id} — {title}",
            "LINE",
            x_axis_title,
            y_axis_title,
            x_units,
            series,
            relevant,
            "Derived from canonical notebook records; blank readings remain gaps.",
        )
        if block:
            blocks.append(block)
    return blocks


def daily_log_plot_blocks(
    tables: dict[str, list[dict[str, Any]]],
) -> list[dict[str, Any]]:
    blocks: list[dict[str, Any]] = []
    daily_rows = []
    for row_number, row in enumerate(tables.get("Daily Log", []), start=2):
        enriched = dict(row)
        enriched["_plot_source_range"] = f"Daily Log row {row_number}"
        daily_rows.append(enriched)
    by_experiment = group_rows(daily_rows, "experiment_id")
    metric_specs = (
        (
            "daily_temperature",
            "logged temperature",
            (("temperature_C", "Temperature", "°C"),),
            "Temperature (°C)",
        ),
        (
            "daily_rpm",
            "logged agitation",
            (("rpm", "RPM", "rpm"),),
            "Agitation speed (rpm)",
        ),
        ("daily_ph", "logged pH", (("pH", "pH", "pH"),), "pH"),
        (
            "daily_particle_size",
            "logged particle size",
            (("particle_size_nm", "Particle size", "nm"),),
            "Particle size (nm)",
        ),
        (
            "daily_viscosity",
            "logged viscosity",
            (("viscosity_cP", "Viscosity", "cP"),),
            "Viscosity (cP)",
        ),
        (
            "daily_percent_outcomes",
            "logged percent outcomes",
            (
                ("solids_percent", "Solids", "%"),
                ("conversion_percent", "Conversion", "%"),
                ("residual_monomer_percent", "Residual monomer", "%"),
            ),
            "Percent (%)",
        ),
        (
            "daily_pdi",
            "logged polydispersity",
            (("polydispersity_index", "PDI", "index"),),
            "Polydispersity index",
        ),
        (
            "daily_tg",
            "logged glass-transition temperature",
            (("Tg_C", "Tg", "°C"),),
            "Temperature (°C)",
        ),
    )
    for experiment_id, rows in by_experiment.items():
        points = daily_log_points(rows)
        blocks.extend(
            metric_blocks(
                experiment_id,
                points,
                metric_specs,
                x_axis_title="Elapsed log time (min)",
                x_units="min",
            )
        )
    return blocks


def daily_log_points(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    parsed = [datetime_value(row.get("timestamp")) for row in rows]
    first = next((value for value in parsed if value is not None), None)
    points: list[dict[str, Any]] = []
    metric_fields = (
        "temperature_C",
        "rpm",
        "pH",
        "solids_percent",
        "particle_size_nm",
        "conversion_percent",
        "viscosity_cP",
        "residual_monomer_percent",
        "polydispersity_index",
        "Tg_C",
    )
    for index, row in enumerate(rows, start=1):
        timestamp = parsed[index - 1]
        elapsed = (
            (timestamp - first).total_seconds() / 60
            if timestamp is not None and first is not None
            else float(index - 1)
        )
        point = {
            "experiment_id": row.get("experiment_id", ""),
            "x_order": index,
            "x_value": elapsed,
            "x_label": row.get("timestamp", "") or str(index),
            "stage": row.get("process_stage", ""),
            "section": "Daily Log",
            "source_record_ids": "",
            "source_ranges": row.get("_plot_source_range", ""),
            "quality_flag": "observed",
            "recorded_at": row.get("timestamp", ""),
        }
        for field in metric_fields:
            point[field] = numeric_value(row.get(field))
        points.append(point)
    return points


def result_trend_plot_blocks(
    tables: dict[str, list[dict[str, Any]]],
) -> list[dict[str, Any]]:
    experiment_dates = {
        str(row.get("experiment_id", "")): str(row.get("date", ""))
        for row in tables.get("Experiments", [])
    }
    grouped: dict[
        tuple[str, str],
        dict[str, Any],
    ] = {}
    for row_number, row in enumerate(tables.get("Results", []), start=2):
        measurement = str(row.get("measurement_type", "")).strip()
        units = str(row.get("units", "")).strip()
        value = numeric_value(row.get("value"))
        if not measurement or value is None:
            continue
        enriched = dict(row)
        enriched["_numeric_value"] = value
        enriched["_plot_source_range"] = f"Results row {row_number}"
        group = grouped.setdefault(
            (measurement.casefold(), units.casefold()),
            {
                "measurement": measurement,
                "units": units,
                "rows": [],
            },
        )
        group["rows"].append(enriched)
    blocks: list[dict[str, Any]] = []
    for _, group in sorted(grouped.items()):
        measurement = str(group["measurement"])
        units = str(group["units"])
        rows = group["rows"]
        rows.sort(
            key=lambda row: (
                experiment_dates.get(str(row.get("experiment_id", "")), ""),
                str(row.get("experiment_id", "")),
                str(row.get("sample_id", "")),
                str(row.get("replicate", "")),
            )
        )
        points = []
        for index, row in enumerate(rows, start=1):
            experiment_id = str(row.get("experiment_id", ""))
            sample_id = str(row.get("sample_id", ""))
            points.append(
                {
                    "experiment_id": experiment_id,
                    "x_order": index,
                    "x_value": float(index),
                    "x_label": " / ".join(
                        value for value in (experiment_id, sample_id) if value
                    ),
                    "result_value": row["_numeric_value"],
                    "stage": row.get("condition", ""),
                    "section": "Results",
                    "source_record_ids": "",
                    "source_ranges": row.get("_plot_source_range", ""),
                    "quality_flag": row.get("quality_flag", "") or "observed",
                    "recorded_at": experiment_dates.get(experiment_id, ""),
                }
            )
        block = make_block(
            "cross_experiment",
            "result_trend",
            f"Cross-experiment trend — {measurement}"
            + (f" ({units})" if units else ""),
            "LINE",
            "Experiment / sample order",
            f"{measurement}" + (f" ({units})" if units else ""),
            "order",
            (
                {
                    "field": "result_value",
                    "name": measurement,
                    "units": units,
                },
            ),
            points,
            "The x-axis is stable experiment/sample order; inspect x_label for identity.",
            identity_suffix=f"{measurement}|{units}",
            x_axis_mode="category",
        )
        if block:
            blocks.append(block)
    return blocks


def make_block(
    experiment_id: str,
    plot_kind: str,
    title: str,
    chart_type: str,
    x_axis_title: str,
    y_axis_title: str,
    x_units: str,
    series: Iterable[dict[str, str]],
    points: list[dict[str, Any]],
    notes: str,
    identity_suffix: str = "",
    x_axis_mode: str | None = None,
) -> dict[str, Any] | None:
    if not points:
        return None
    series_tuple = tuple(series)
    plot_id = stable_id("plot", experiment_id, plot_kind, identity_suffix)
    return {
        "plot_id": plot_id,
        "experiment_id": experiment_id,
        "plot_kind": plot_kind,
        "title": title,
        "chart_type": chart_type,
        "x_axis_title": x_axis_title,
        "y_axis_title": y_axis_title,
        "x_units": x_units,
        "series": series_tuple,
        "points": points,
        "notes": notes,
        "x_axis_mode": x_axis_mode
        or ("category" if chart_type.upper() == "COLUMN" else "numeric"),
    }


def group_rows(
    rows: Iterable[dict[str, Any]],
    field: str,
) -> dict[str, list[dict[str, Any]]]:
    grouped: dict[str, list[dict[str, Any]]] = {}
    for row in rows:
        if not isinstance(row, dict):
            continue
        key = str(row.get(field, "")).strip()
        if key:
            grouped.setdefault(key, []).append(row)
    return grouped


def clock_minute(value: Any) -> float | None:
    if isinstance(value, datetime):
        return value.hour * 60 + value.minute + value.second / 60
    if isinstance(value, time):
        return value.hour * 60 + value.minute + value.second / 60
    text = str(value or "").strip()
    for fmt in ("%I:%M:%S %p", "%I:%M %p", "%H:%M:%S", "%H:%M"):
        try:
            parsed = datetime.strptime(text, fmt)
            return parsed.hour * 60 + parsed.minute + parsed.second / 60
        except ValueError:
            continue
    parsed_datetime = datetime_value(value)
    if parsed_datetime is not None:
        return (
            parsed_datetime.hour * 60
            + parsed_datetime.minute
            + parsed_datetime.second / 60
        )
    return None


def datetime_value(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, date):
        return datetime.combine(value, time())
    text = str(value or "").strip()
    if not text:
        return None
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).replace(
            tzinfo=None
        )
    except ValueError:
        pass
    for fmt in (
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%m/%d/%Y %H:%M",
        "%m/%d/%Y %I:%M %p",
    ):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def apply_plot_report_to_workbook(
    workbook_path: str | Path,
    report: dict[str, Any],
    output_workbook: str | Path | None = None,
) -> Path:
    source = Path(workbook_path).expanduser().resolve()
    destination = (
        Path(output_workbook).expanduser().resolve() if output_workbook else source
    )
    workbook = load_workbook(source)
    ensure_workbook_contract(workbook)
    desired = {
        PLOT_DATA_SHEET: report.get("plot_data_rows", []),
        PLOT_DEFINITIONS_SHEET: report.get("plot_definitions", []),
        PLOT_DASHBOARD_SHEET: report.get("plot_dashboard_rows", []),
    }
    for sheet_name, rows in desired.items():
        worksheet = workbook[sheet_name]
        if worksheet.max_row > 1:
            worksheet.delete_rows(2, worksheet.max_row - 1)
        headers = [str(cell.value or "") for cell in worksheet[1]]
        for row in rows:
            worksheet.append([row.get(header, "") for header in headers])
        worksheet.freeze_panes = "A2"

    data_sheet = workbook[PLOT_DATA_SHEET]
    data_headers = [str(cell.value or "") for cell in data_sheet[1]]
    for row_number in range(2, data_sheet.max_row + 1):
        if data_sheet.cell(row_number, 1).value == "header":
            for column_number in range(1, len(data_headers) + 1):
                cell = data_sheet.cell(row_number, column_number)
                cell.fill = PatternFill("solid", fgColor="D9EAF7")
                cell.font = Font(bold=True)

    dashboard = workbook[PLOT_DASHBOARD_SHEET]
    dashboard._charts = []
    ready_definitions = [
        definition
        for definition in report.get("plot_definitions", [])
        if definition.get("status") == "ready"
    ]
    for index, definition in enumerate(ready_definitions):
        chart = excel_chart_for_definition(data_sheet, definition)
        column = "I" if index % 2 == 0 else "T"
        row = 2 + (index // 2) * 20
        dashboard.add_chart(chart, f"{column}{row}")
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    return destination


def excel_chart_for_definition(data_sheet: Any, definition: dict[str, Any]) -> Any:
    chart_type = str(definition.get("chart_type", "LINE")).upper()
    if chart_type == "COLUMN":
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
    else:
        chart = LineChart()
        chart.style = 13
        chart.smooth = False
    chart.title = str(definition.get("title", ""))
    chart.x_axis.title = str(definition.get("x_axis_title", ""))
    chart.y_axis.title = str(definition.get("y_axis_title", ""))
    chart.height = 8
    chart.width = 15
    start_row = int(definition["data_start_row"])
    end_row = int(definition["data_end_row"])
    data_headers = list(sheet_by_name(PLOT_DATA_SHEET).headers)
    x_field = (
        "x_label"
        if chart_type == "COLUMN"
        or definition.get("x_axis_mode") == "category"
        else "x_value"
    )
    x_column = data_headers.index(x_field) + 1
    categories = Reference(
        data_sheet,
        min_col=x_column,
        min_row=start_row + 1,
        max_row=end_row,
    )
    chart.set_categories(categories)
    for series_number, value_field in enumerate(SERIES_VALUE_COLUMNS, start=1):
        if not definition.get(f"series_{series_number}_name"):
            continue
        value_column = data_headers.index(value_field) + 1
        data = Reference(
            data_sheet,
            min_col=value_column,
            min_row=start_row,
            max_row=end_row,
        )
        chart.add_data(data, titles_from_data=True)
    return chart


def google_chart_requests(
    report: dict[str, Any],
    sheet_ids: dict[str, int],
    metadata: dict[str, Any],
) -> list[dict[str, Any]]:
    if PLOT_DATA_SHEET not in sheet_ids or PLOT_DASHBOARD_SHEET not in sheet_ids:
        return []
    existing: dict[str, list[dict[str, Any]]] = {}
    for sheet in metadata.get("sheets", []) or []:
        if not isinstance(sheet, dict):
            continue
        for chart in sheet.get("charts", []) or []:
            if not isinstance(chart, dict):
                continue
            title = str((chart.get("spec") or {}).get("title", ""))
            match = MANAGED_CHART_TITLE_RE.search(title)
            if match:
                existing.setdefault(match.group(1), []).append(chart)

    requests: list[dict[str, Any]] = []
    desired_ids: set[str] = set()
    ready_definitions = [
        definition
        for definition in report.get("plot_definitions", [])
        if definition.get("status") == "ready"
    ]
    for index, definition in enumerate(ready_definitions):
        plot_id = str(definition["plot_id"])
        desired_ids.add(plot_id)
        spec = google_chart_spec(definition, sheet_ids[PLOT_DATA_SHEET])
        position = google_chart_position(sheet_ids[PLOT_DASHBOARD_SHEET], index)
        matches = existing.get(plot_id, [])
        if matches:
            chart_id = int(matches[0]["chartId"])
            requests.append(
                {"updateChartSpec": {"chartId": chart_id, "spec": spec}}
            )
            requests.append(
                {
                    "updateEmbeddedObjectPosition": {
                        "objectId": chart_id,
                        "newPosition": position,
                        "fields": "overlayPosition",
                    }
                }
            )
            for duplicate in matches[1:]:
                requests.append(
                    {
                        "deleteEmbeddedObject": {
                            "objectId": int(duplicate["chartId"])
                        }
                    }
                )
        else:
            requests.append(
                {"addChart": {"chart": {"spec": spec, "position": position}}}
            )
    for plot_id, charts in existing.items():
        if plot_id in desired_ids:
            continue
        for chart in charts:
            requests.append(
                {
                    "deleteEmbeddedObject": {
                        "objectId": int(chart["chartId"])
                    }
                }
            )
    return requests


def google_chart_spec(
    definition: dict[str, Any],
    plot_data_sheet_id: int,
) -> dict[str, Any]:
    chart_type = str(definition.get("chart_type", "LINE")).upper()
    headers = list(sheet_by_name(PLOT_DATA_SHEET).headers)
    start_row_index = int(definition["data_start_row"]) - 1
    end_row_index = int(definition["data_end_row"])
    x_field = (
        "x_label"
        if chart_type == "COLUMN"
        or definition.get("x_axis_mode") == "category"
        else "x_value"
    )

    def source_range(field: str) -> dict[str, Any]:
        column_index = headers.index(field)
        return {
            "sources": [
                {
                    "sheetId": plot_data_sheet_id,
                    "startRowIndex": start_row_index,
                    "endRowIndex": end_row_index,
                    "startColumnIndex": column_index,
                    "endColumnIndex": column_index + 1,
                }
            ]
        }

    series = []
    for series_number, field in enumerate(SERIES_VALUE_COLUMNS, start=1):
        if definition.get(f"series_{series_number}_name"):
            series.append(
                {
                    "series": {"sourceRange": source_range(field)},
                    "targetAxis": "LEFT_AXIS",
                }
            )
    title = (
        f"{definition.get('title', '')} "
        f"[LNA:{definition.get('plot_id', '')}]"
    )
    basic_chart = {
        "chartType": chart_type,
        "legendPosition": "BOTTOM_LEGEND",
        "axis": [
            {
                "position": "BOTTOM_AXIS",
                "title": str(definition.get("x_axis_title", "")),
            },
            {
                "position": "LEFT_AXIS",
                "title": str(definition.get("y_axis_title", "")),
            },
        ],
        "domains": [
            {"domain": {"sourceRange": source_range(x_field)}}
        ],
        "series": series,
        "headerCount": 1,
        "interpolateNulls": False,
    }
    if chart_type in {"BAR", "COLUMN", "AREA"}:
        basic_chart["stackedType"] = "NOT_STACKED"
    return {
        "title": title,
        "basicChart": basic_chart,
    }


def google_chart_position(
    dashboard_sheet_id: int,
    index: int,
) -> dict[str, Any]:
    return {
        "overlayPosition": {
            "anchorCell": {
                "sheetId": dashboard_sheet_id,
                "rowIndex": 1 + (index // 2) * 20,
                "columnIndex": 8 + (index % 2) * 10,
            },
            "offsetXPixels": 8,
            "offsetYPixels": 8,
            "widthPixels": 600,
            "heightPixels": 360,
        }
    }
