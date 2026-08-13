from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.views import Selection

from .batch_builder import BATCH_BUILDER_FORMULA_COLUMNS, excel_batch_builder_formula
from .reaction_outcomes import (
    REACTION_OUTCOME_FORMULA_COLUMNS,
    excel_reaction_outcome_formula,
)
from .schema import (
    CONTROLLED_VOCAB_VALIDATIONS,
    RUN_CONSOLE_SHEET,
    SHEETS,
    column_number_format,
)
from .scientist_workspace import (
    PROCESS_METRICS,
    reaction_master_excel_formula,
)


HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SUBTLE_FILL = PatternFill("solid", fgColor="D9EAF7")
CONSOLE_NAVY = "122E4A"
CONSOLE_TEAL = "197A8C"
CONSOLE_PALE_BLUE = "EAF4F7"
CONSOLE_INPUT = "FFF2B8"
TECHNICAL_SHEETS = frozenset(
    {
        "Daily Reviews",
        "Project Notebook Records",
        "Source Sync",
        "Plot Data",
        "Plot Definitions",
        "Process Knowledge",
        "Controlled Vocab",
        "Agent Config",
        "Workbook Metadata",
        "Audit Log",
        "Formulations",
        "Daily Log",
        "Results",
        "Plot Dashboard",
    }
)
CORE_ENTRY_SHEETS = frozenset(
    {
        "Experiments",
        "Reaction Master",
        "Batch Builder",
        "Bench Log",
        "Measurements",
        "Plot Studio",
        "Run Capture Plan",
        "Samples",
        "Deviations",
        "Raw Data Files",
        "Inventory Transactions",
        "Equipment Bookings",
        "Record Signatures",
        "Notebook Sections",
        "Reaction Outcomes",
    }
)
REFERENCE_SHEETS = frozenset(
    {"Master Reagents", "Equipment", "Protocols", "Specifications", "Experiment Templates"}
)


def freeze_pane_for_sheet(sheet_name: str) -> str:
    if sheet_name in {
        "Daily Log",
        "Bench Log",
        "Batch Builder",
        "Formulations",
        "Results",
        "Measurements",
        "Run Capture Plan",
        "Samples",
        "Deviations",
        "Raw Data Files",
        "Inventory Transactions",
        "Equipment Bookings",
        "Record Signatures",
        "Notebook Sections",
        "Reaction Outcomes",
        "Project Notebook Records",
    }:
        return "C2"
    return "B2"


def column_width(header: str) -> float:
    text = header.lower()
    if any(
        token in text
        for token in (
            "json",
            "notes",
            "description",
            "objective",
            "hypothesis",
            "summary",
            "finding",
            "rationale",
            "proposed_change",
            "expected_effect",
            "interpretation",
            "guidance",
            "immediate_action",
            "impact_assessment",
            "disposition",
            "change_summary",
        )
    ):
        return 40
    if "url" in text or "source_range" in text or "linked_" in text:
        return 28
    if text in {
        "title",
        "name",
        "material_name",
        "reagent_name",
        "measurement_type",
        "parameter",
        "action",
        "method",
        "search_terms",
        "typical_examples",
        "measured_fields",
    }:
        return 26
    if text.endswith("_at") or text.endswith("_date") or text == "timestamp":
        return 22
    if text.endswith("_id") or text in {"record_id", "source_key"}:
        return 20
    if text in {
        "status",
        "qc_status",
        "quality_flag",
        "process_stage",
        "category",
        "target_role",
        "units",
        "confidence",
        "active",
        "required",
    }:
        return 16
    return 18


def sheet_tab_color(sheet_name: str) -> str:
    if sheet_name == RUN_CONSOLE_SHEET:
        return CONSOLE_NAVY
    if sheet_name in CORE_ENTRY_SHEETS:
        return "197A8C"
    if sheet_name in REFERENCE_SHEETS:
        return "4F8F5B"
    if sheet_name in {"Plot Dashboard", "Plot Studio", "Literature Evidence"}:
        return "6650A3"
    if sheet_name == "Agent Suggestions":
        return "D1842B"
    return "8C999F"


def console_lookup(column_letter: str) -> str:
    return (
        '=IF($B$3="","",IFERROR(IF(INDEX(Experiments!$'
        f'{column_letter}$2:${column_letter}$1000,'
        'MATCH($B$3,Experiments!$A$2:$A$1000,0))="","Not recorded",'
        'INDEX(Experiments!$'
        f'{column_letter}$2:${column_letter}$1000,'
        'MATCH($B$3,Experiments!$A$2:$A$1000,0))),"Not recorded"))'
    )


def local_active_run_queue_formula() -> str:
    return '=IF($D30="","",IF($D30="running","1 · RUNNING",IF($D30="planned","2 · PLANNED","3 · REVIEW")))'


def active_run_source_index(queue_row: int) -> str:
    return (
        'AGGREGATE(15,6,(ROW(Experiments!$A$2:$A$1000)-ROW(Experiments!$A$2)+1)/'
        '(((Experiments!$I$2:$I$1000="running")+(Experiments!$I$2:$I$1000="planned")+'
        '((Experiments!$I$2:$I$1000="needs_review")*(Experiments!$L$2:$L$1000="")))>0),'
        f'ROWS($A$30:A{queue_row}))'
    )


def excel_plot_row_index(source_sheet: str, value_expression: str, output_row: int, anchor_column: str) -> str:
    return (
        f'AGGREGATE(15,6,(ROW(\'{source_sheet}\'!$A$2:$A$1000)-ROW(\'{source_sheet}\'!$A$2)+1)/'
        f'((\'{source_sheet}\'!$A$2:$A$1000=$B$3)*({value_expression}<>"")),'
        f'ROWS(${anchor_column}$9:{anchor_column}{output_row}))'
    )


def ensure_run_console(workbook: Workbook) -> None:
    is_new = RUN_CONSOLE_SHEET not in workbook.sheetnames
    selected_experiment = ""
    if is_new:
        worksheet = workbook.create_sheet(RUN_CONSOLE_SHEET, 0)
    else:
        worksheet = workbook[RUN_CONSOLE_SHEET]
        selected_experiment = worksheet["B3"].value or ""
        if workbook.sheetnames[0] != RUN_CONSOLE_SHEET:
            workbook._sheets.remove(worksheet)
            workbook._sheets.insert(0, worksheet)

    for row in worksheet.iter_rows(min_row=1, max_row=40, min_col=1, max_col=10):
        for cell in row:
            if cell.coordinate != "B3":
                cell.value = None

    worksheet["A1"] = "COCHRAN LAB NOTEBOOK"
    worksheet["A2"] = (
        "Scientist run console · select an experiment to inspect readiness and "
        "navigate the record"
    )
    worksheet["A3"] = "Active experiment"
    worksheet["B3"] = selected_experiment

    for cell, value in {
        "A5": "EXPERIMENT OVERVIEW",
        "E5": "READINESS",
        "F5": "VALUE",
        "G5": "CHECK",
        "I5": "QUICK LINKS",
        "A6": "Status",
        "B6": console_lookup("I"),
        "A7": "Operator",
        "B7": console_lookup("H"),
        "A8": "Process",
        "B8": console_lookup("D"),
        "A9": "Date",
        "B9": console_lookup("B"),
        "A10": "Objective",
        "B10": console_lookup("E"),
        "A11": "Hypothesis",
        "B11": console_lookup("F"),
        "A12": "Next step",
        "B12": console_lookup("J"),
        "E6": "Operator",
        "F6": "=$B$7",
        "G6": '=IF($B$3="","",IF(F6<>"Not recorded","✓ Ready","⚠ Missing"))',
        "E7": "Protocol",
        "F7": console_lookup("Q"),
        "G7": '=IF($B$3="","",IF(F7<>"Not recorded","✓ Ready","⚠ Missing"))',
        "E8": "Equipment",
        "F8": console_lookup("R"),
        "G8": '=IF($B$3="","",IF(F8<>"Not recorded","✓ Ready","⚠ Missing"))',
        "E9": "Batch charges",
        "F9": '=IF($B$3="","",COUNTIF(\'Batch Builder\'!$A$2:$A$1000,$B$3))',
        "G9": '=IF($B$3="","",IF(F9=0,"⚠ Missing",IF(COUNTIFS(\'Batch Builder\'!$A$2:$A$1000,$B$3,\'Batch Builder\'!$AR$2:$AR$1000,"READY")=F9,"✓ Quantified","⚠ Fix mass plan")))',
        "E10": "Run steps",
        "F10": '=IF($B$3="","",COUNTIF(\'Run Capture Plan\'!$A$2:$A$1000,$B$3))',
        "G10": '=IF($B$3="","",IF(F10>0,"✓ Ready","⚠ Missing"))',
        "E11": "Observations",
        "F11": '=IF($B$3="","",COUNTIF(\'Bench Log\'!$A$2:$A$1000,$B$3)+COUNTIF(\'Daily Log\'!$A$2:$A$1000,$B$3))',
        "G11": '=IF($B$3="","",IF(F11>0,"✓ Logged","⚠ Missing"))',
        "E12": "Samples",
        "F12": '=IF($B$3="","",COUNTIF(Samples!$B$2:$B$1000,$B$3))',
        "G12": '=IF($B$3="","",IF(F12>0,"✓ Logged","⚠ Missing"))',
        "E13": "Results",
        "F13": '=IF($B$3="","",COUNTIF(Measurements!$A$2:$A$1000,$B$3)+COUNTIF(Results!$A$2:$A$1000,$B$3))',
        "G13": '=IF($B$3="","",IF(F13>0,"✓ Logged","⚠ Missing"))',
        "E14": "Raw files",
        "F14": '=IF($B$3="","",COUNTIF(\'Raw Data Files\'!$B$2:$B$1000,$B$3))',
        "G14": '=IF($B$3="","",IF(F14>0,"✓ Linked","⚠ Missing"))',
        "E15": "Open deviations",
        "F15": '=IF($B$3="","",COUNTIFS(Deviations!$B$2:$B$1000,$B$3,Deviations!$K$2:$K$1000,"<>closed"))',
        "G15": '=IF($B$3="","",IF(F15=0,"✓ Clear","⚠ Attention"))',
        "E16": "Reviewer",
        "F16": console_lookup("U"),
        "G16": '=IF($B$3="","",IF(F16<>"Not recorded","✓ Ready","⚠ Missing"))',
        "E17": "Completeness",
        "F17": '=IF($B$3="","",COUNTIF($G$6:$G$16,"✓*")/11)',
        "G17": '=IF($B$3="","",IF(F17=1,"✓ Ready to close",TEXT(F17,"0%")&" complete"))',
        "A18": "BENCH WORKFLOW",
        "A19": "1 · PLAN",
        "B19": "Set protocol, equipment, run steps, acceptance criteria, and sample plan before starting.",
        "A20": "2 · PREPARE",
        "B20": "Confirm reagent lots, equipment calibration, formulation targets, and safety controls.",
        "A21": "3 · RUN",
        "B21": "Record timestamps, actual additions, process conditions, observations, and deviations as they happen.",
        "A22": "4 · MEASURE",
        "B22": "Create sample records, link instrument files, capture uncertainty, and evaluate specifications.",
        "A23": "5 · REVIEW",
        "B23": "Resolve deviations, document the conclusion, then sign and independently witness the finalized record when required.",
        "A25": "GOOD RECORDS",
        "B25": "Use stable IDs. Record actual values, times, lots, and operators. Link raw evidence; never replace it with a summary.",
        "A28": "ACTIVE RUN QUEUE",
        "A29": "PRIORITY",
        "B29": "EXPERIMENT",
        "C29": "DATE",
        "D29": "STATUS",
        "E29": "OWNER",
        "F29": "NEXT ACTION",
        "G29": "COMPLETE",
        "H29": "RECORD",
        "A30": local_active_run_queue_formula(),
        "I28": "QUEUE SUMMARY",
        "J28": "COUNT",
        "I29": "Current queue",
        "J29": '=COUNTIFS(Experiments!$I$2:$I$1000,"running")+COUNTIFS(Experiments!$I$2:$I$1000,"planned")+COUNTIFS(Experiments!$I$2:$I$1000,"needs_review",Experiments!$L$2:$L$1000,"")',
        "I30": "Running",
        "J30": '=COUNTIF(Experiments!$I$2:$I$1000,"running")',
        "I31": "Planned",
        "J31": '=COUNTIF(Experiments!$I$2:$I$1000,"planned")',
        "I32": "Current review",
        "J32": '=COUNTIFS(Experiments!$I$2:$I$1000,"needs_review",Experiments!$L$2:$L$1000,"")',
        "I33": "Imported history",
        "J33": '=COUNTIF(Experiments!$L$2:$L$1000,"<>")',
        "I34": "History needing review",
        "J34": '=COUNTIFS(Experiments!$L$2:$L$1000,"<>",Experiments!$I$2:$I$1000,"needs_review")',
    }.items():
        worksheet[cell] = value

    for queue_row in range(30, 50):
        source_index = active_run_source_index(queue_row)
        worksheet[f"B{queue_row}"] = f'=IFERROR(INDEX(Experiments!$A$2:$A$1000,{source_index}),"")'
        worksheet[f"C{queue_row}"] = f'=IF($B{queue_row}="","",INDEX(Experiments!$B$2:$B$1000,{source_index}))'
        worksheet[f"D{queue_row}"] = f'=IF($B{queue_row}="","",INDEX(Experiments!$I$2:$I$1000,{source_index}))'
        worksheet[f"E{queue_row}"] = f'=IF($B{queue_row}="","",IF(INDEX(Experiments!$H$2:$H$1000,{source_index})="","Unassigned",INDEX(Experiments!$H$2:$H$1000,{source_index})))'
        worksheet[f"A{queue_row}"] = f'=IF($D{queue_row}="","",IF($D{queue_row}="running","1 · RUNNING",IF($D{queue_row}="planned","2 · PLANNED","3 · REVIEW")))'
        worksheet[f"F{queue_row}"] = (
            f'=IF($B{queue_row}="","",IF($E{queue_row}="Unassigned","Assign operator",'
            f'IF(COUNTIF(\'Batch Builder\'!$A$2:$A$1000,$B{queue_row})=0,"Enter batch quantities",'
            f'IF(COUNTIF(\'Run Capture Plan\'!$A$2:$A$1000,$B{queue_row})=0,"Build run plan",'
            f'IF(COUNTIF(\'Raw Data Files\'!$B$2:$B$1000,$B{queue_row})=0,"Link raw file","Review record")))))'
        )
        worksheet[f"G{queue_row}"] = (
            f'=IF($B{queue_row}="","",(N($E{queue_row}<>"Unassigned")+'
            f'N(COUNTIF(\'Batch Builder\'!$A$2:$A$1000,$B{queue_row})>0)+'
            f'N(COUNTIF(\'Run Capture Plan\'!$A$2:$A$1000,$B{queue_row})>0)+'
            f'N((COUNTIF(\'Bench Log\'!$A$2:$A$1000,$B{queue_row})+COUNTIF(\'Daily Log\'!$A$2:$A$1000,$B{queue_row}))>0)+'
            f'N(COUNTIF(Samples!$B$2:$B$1000,$B{queue_row})>0)+'
            f'N((COUNTIF(Measurements!$A$2:$A$1000,$B{queue_row})+COUNTIF(Results!$A$2:$A$1000,$B{queue_row}))>0)+'
            f'N(COUNTIF(\'Raw Data Files\'!$B$2:$B$1000,$B{queue_row})>0))/7)'
        )
        worksheet[f"G{queue_row}"].number_format = "0%"
        worksheet[f"H{queue_row}"] = f'=IF($B{queue_row}="","",HYPERLINK("#\'Experiments\'!A"&MATCH($B{queue_row},Experiments!$A$2:$A$1000,0)+1,"Open →"))'

    for row_number, (label, sheet_name) in enumerate(
        (
            ("Experiment record", "Experiments"),
            ("Run plan", "Run Capture Plan"),
            ("Bench log", "Bench Log"),
            ("Batch quantities", "Batch Builder"),
            ("Samples", "Samples"),
            ("Measurements", "Measurements"),
            ("Raw files", "Raw Data Files"),
            ("Deviations", "Deviations"),
            ("Plot studio", "Plot Studio"),
            ("Material ledger", "Inventory Transactions"),
            ("Equipment schedule", "Equipment Bookings"),
            ("Signatures", "Record Signatures"),
        ),
        start=6,
    ):
        worksheet.cell(row=row_number, column=9, value=label)
        worksheet.cell(
            row=row_number,
            column=10,
            value=f'=HYPERLINK("#\'{sheet_name}\'!A1","Open →")',
        )

    worksheet.freeze_panes = "A4"
    worksheet.sheet_view.showGridLines = False
    worksheet.sheet_properties.tabColor = sheet_tab_color(RUN_CONSOLE_SHEET)
    worksheet.column_dimensions["A"].width = 20
    worksheet.column_dimensions["B"].width = 35
    worksheet.column_dimensions["C"].width = 12
    worksheet.column_dimensions["D"].width = 12
    worksheet.column_dimensions["E"].width = 18
    worksheet.column_dimensions["F"].width = 25
    worksheet.column_dimensions["G"].width = 13
    worksheet.column_dimensions["H"].width = 12
    worksheet.column_dimensions["I"].width = 18
    worksheet.column_dimensions["J"].width = 12
    for row_number, height in ((1, 34), (2, 22), (3, 28), (5, 24), (18, 24), (25, 38), (28, 24), (29, 22)):
        worksheet.row_dimensions[row_number].height = height

    thin_gold = Side(style="thin", color="BD9729")
    for row in worksheet.iter_rows(min_row=1, max_row=40, min_col=1, max_col=10):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="1F2933")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.fill = PatternFill("solid", fgColor="FFFFFF")
    for cell in worksheet[1][:10]:
        cell.fill = PatternFill("solid", fgColor=CONSOLE_NAVY)
        cell.font = Font(name="Arial", size=20, bold=True, color="FFFFFF")
    for cell in worksheet[2][:10]:
        cell.fill = PatternFill("solid", fgColor="DDEBF2")
        cell.font = Font(name="Arial", size=10, color="384F5E")
    worksheet["B3"].fill = PatternFill("solid", fgColor=CONSOLE_INPUT)
    worksheet["B3"].font = Font(name="Arial", size=10, bold=True, color="1F2933")
    worksheet["B3"].border = Border(
        left=thin_gold, right=thin_gold, top=thin_gold, bottom=thin_gold
    )
    for row_number in (5, 18, 28):
        for cell in worksheet[row_number][:10]:
            cell.fill = PatternFill("solid", fgColor=CONSOLE_TEAL)
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    for row_number in range(6, 18):
        for column in (1, 5):
            cell = worksheet.cell(row=row_number, column=column)
            cell.fill = PatternFill("solid", fgColor=CONSOLE_PALE_BLUE)
            cell.font = Font(name="Arial", size=10, bold=True, color="1F2933")
    for row_number in range(19, 24):
        cell = worksheet.cell(row=row_number, column=1)
        cell.fill = PatternFill("solid", fgColor="DDEBF2")
        cell.font = Font(name="Arial", size=10, bold=True, color="1F2933")
    for cell in worksheet[25][:10]:
        cell.fill = PatternFill("solid", fgColor=CONSOLE_INPUT)
        cell.font = Font(name="Arial", size=10, bold=True, color="1F2933")
    worksheet["F17"].number_format = "0%"
    worksheet["B9"].number_format = "yyyy-mm-dd"
    for cell in worksheet[29][:8]:
        cell.fill = PatternFill("solid", fgColor="DDEBF2")
        cell.font = Font(name="Arial", size=10, bold=True, color="1F2933")
    for cell in worksheet[28][8:10]:
        cell.fill = PatternFill("solid", fgColor=CONSOLE_TEAL)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    for row_number in range(29, 35):
        worksheet.cell(row=row_number, column=9).fill = PatternFill(
            "solid", fgColor=CONSOLE_PALE_BLUE
        )
        worksheet.cell(row=row_number, column=9).font = Font(
            name="Arial", size=10, bold=True, color="1F2933"
        )
    for row_number in range(30, 41):
        worksheet.cell(row=row_number, column=3).number_format = "yyyy-mm-dd"
        worksheet.cell(row=row_number, column=7).number_format = "0%"

    has_validation = any(
        "B3" in str(validation.sqref)
        for validation in worksheet.data_validations.dataValidation
    )
    if not has_validation:
        validation = DataValidation(
            type="list",
            formula1="'Experiments'!$A$2:$A$1000",
            allow_blank=True,
        )
        validation.error = "Choose an experiment ID from the Experiments table."
        validation.errorTitle = "Unknown experiment"
        validation.prompt = "Select the experiment to inspect."
        validation.promptTitle = "Active experiment"
        worksheet.add_data_validation(validation)
        validation.add("B3")
    if is_new:
        worksheet.conditional_formatting.add(
            "G6:G17",
            FormulaRule(
                formula=['LEFT(G6,1)="✓"'],
                fill=PatternFill("solid", fgColor="C6E0B4"),
                font=Font(bold=True),
            ),
        )
        worksheet.conditional_formatting.add(
            "G6:G17",
            FormulaRule(
                formula=['LEFT(G6,1)="⚠"'],
                fill=PatternFill("solid", fgColor="FFE699"),
                font=Font(bold=True),
            ),
        )


def apply_workbook_presentation(workbook: Workbook) -> None:
    ensure_run_console(workbook)
    for spec in SHEETS:
        if spec.name not in workbook.sheetnames:
            continue
        worksheet = workbook[spec.name]
        worksheet.sheet_properties.tabColor = sheet_tab_color(spec.name)
        worksheet.sheet_state = "hidden" if spec.name in TECHNICAL_SHEETS else "visible"
        worksheet.sheet_view.showGridLines = spec.name != "Plot Dashboard"
        worksheet.freeze_panes = freeze_pane_for_sheet(spec.name)
        worksheet.row_dimensions[1].height = 27
        for index, header in enumerate(spec.headers, start=1):
            worksheet.column_dimensions[get_column_letter(index)].width = column_width(header)
        for cell in worksheet[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="center")
    ensure_batch_builder(workbook)
    ensure_reaction_outcomes(workbook)
    ensure_reaction_master(workbook)
    ensure_plot_studio(workbook)


def ensure_reaction_master(workbook: Workbook, end_row: int = 1000) -> None:
    """Populate the visible formula-driven index of every reaction run."""

    if "Reaction Master" not in workbook.sheetnames:
        return
    worksheet = workbook["Reaction Master"]
    headers = [str(cell.value or "") for cell in worksheet[1]]
    for row_number in range(2, end_row + 1):
        for column_number, header in enumerate(headers, start=1):
            cell = worksheet.cell(row=row_number, column=column_number)
            cell.value = reaction_master_excel_formula(header, row_number)
            cell.fill = PatternFill("solid", fgColor=CONSOLE_PALE_BLUE)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    worksheet.freeze_panes = "B2"
    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{end_row}"
    for column in ("B", "P", "T", "AQ", "AS"):
        for row_number in range(2, end_row + 1):
            worksheet[f"{column}{row_number}"].number_format = "yyyy-mm-dd hh:mm"
    for column in ("J", "K", "L", "U", "V", "W", "AD", "AE", "AF", "AG", "AH", "AI", "AJ"):
        for row_number in range(2, end_row + 1):
            worksheet[f"{column}{row_number}"].number_format = "0.00"
    for row_number in range(2, end_row + 1):
        worksheet[f"X{row_number}"].number_format = "0%"
        worksheet[f"Y{row_number}"].number_format = "0"
        worksheet[f"Z{row_number}"].number_format = "0"
        worksheet[f"AN{row_number}"].number_format = "0"


def ensure_batch_builder(workbook: Workbook, end_row: int = 1000) -> None:
    """Install scientist-facing quantity formulas and visual input cues."""

    if "Batch Builder" not in workbook.sheetnames:
        return
    worksheet = workbook["Batch Builder"]
    headers = [str(cell.value or "") for cell in worksheet[1]]
    formula_columns = {
        header: headers.index(header) + 1 for header in BATCH_BUILDER_FORMULA_COLUMNS
    }
    for row_number in range(2, max(2, end_row) + 1):
        for header, column_number in formula_columns.items():
            cell = worksheet.cell(row=row_number, column=column_number)
            cell.value = excel_batch_builder_formula(header, row_number)
            cell.fill = PatternFill("solid", fgColor=CONSOLE_PALE_BLUE)
            cell.font = Font(name="Arial", size=10, color="1F2933")
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    input_headers = (
        "experiment_id",
        "charge_id",
        "stage",
        "charge_type",
        "target_role",
        "feed_order",
        "reagent_id",
        "parts_per_hundred_monomer",
        "stage_monomer_basis_g",
        "direct_mass_g",
        "stock_active_fraction",
        "density_override_g_mL",
        "actual_mass_g",
        "feed_start_min",
        "feed_duration_min",
        "target_temperature_C",
        "lot",
        "recorded_by",
        "recorded_at",
        "charge_status",
        "notes",
        "equivalent_basis_mmol",
        "calculation_mode",
        "recipe_wt_percent",
        "target_active_mass_g",
        "target_equivalents",
        "mass_tolerance_percent",
        "target_functional_equivalents",
        "weighing_method",
        "source_container_before_g",
        "source_container_after_g",
        "carrier_reagent_id",
    )
    for header in input_headers:
        column_number = headers.index(header) + 1
        for row_number in range(2, max(2, end_row) + 1):
            worksheet.cell(row=row_number, column=column_number).fill = PatternFill(
                "solid", fgColor="FFF9E3"
            )
    worksheet.freeze_panes = "C2"
    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(2, end_row)}"
    worksheet.row_dimensions[1].height = 42
    worksheet.conditional_formatting.add(
        f"AP2:AP{end_row}",
        FormulaRule(
            formula=['AND($A2<>"",$AP2<>"",IF($AN2<>"",$AN2,$AO2)>0,ABS($AP2)>IF($AN2<>"",$AN2,$AO2))'],
            fill=PatternFill("solid", fgColor="F4CCCC"),
            font=Font(bold=True),
        ),
    )
    for status, color in (("PASS", "C6E0B4"), ("FAIL", "F4CCCC")):
        worksheet.conditional_formatting.add(
            f"AQ2:AQ{end_row}",
            CellIsRule(
                operator="equal",
                formula=[f'"{status}"'],
                fill=PatternFill("solid", fgColor=color),
                font=Font(bold=True),
            ),
        )
    worksheet.conditional_formatting.add(
        f"AR2:AR{end_row}",
        FormulaRule(
            formula=['AND($A2<>"",$AR2<>"READY")'],
            fill=PatternFill("solid", fgColor="FFE699"),
            font=Font(bold=True),
        ),
    )
    worksheet.conditional_formatting.add(
        f"BF2:BF{end_row}",
        FormulaRule(
            formula=['AND($A2<>"",$BF2<>"READY",$BF2<>"NOT_RECORDED")'],
            fill=PatternFill("solid", fgColor="FFE699"),
            font=Font(bold=True),
        ),
    )


def ensure_reaction_outcomes(workbook: Workbook, end_row: int = 1000) -> None:
    """Install yield and material-balance formulas while preserving clear inputs."""

    if "Reaction Outcomes" not in workbook.sheetnames:
        return
    worksheet = workbook["Reaction Outcomes"]
    headers = [str(cell.value or "") for cell in worksheet[1]]
    calculated = set(REACTION_OUTCOME_FORMULA_COLUMNS)
    for row_number in range(2, max(2, end_row) + 1):
        for header in calculated:
            column_number = headers.index(header) + 1
            cell = worksheet.cell(row=row_number, column=column_number)
            cell.value = excel_reaction_outcome_formula(header, row_number)
            cell.fill = PatternFill("solid", fgColor=CONSOLE_PALE_BLUE)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for header in set(headers) - calculated:
            worksheet.cell(row=row_number, column=headers.index(header) + 1).fill = PatternFill(
                "solid", fgColor="FFF9E3"
            )
    worksheet.freeze_panes = "C2"
    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(2, end_row)}"
    for header in ("isolated_yield_percent", "material_balance_closure_percent"):
        column = get_column_letter(headers.index(header) + 1)
        for row_number in range(2, end_row + 1):
            worksheet[f"{column}{row_number}"].number_format = "0.0%"
    status_column = get_column_letter(headers.index("mass_balance_status") + 1)
    for status, color in (("CLOSED", "C6E0B4"), ("OPEN", "FFE699"), ("OVER_ACCOUNTED", "F4CCCC")):
        worksheet.conditional_formatting.add(
            f"{status_column}2:{status_column}{end_row}",
            CellIsRule(operator="equal", formula=[f'"{status}"'], fill=PatternFill("solid", fgColor=color)),
        )


def ensure_plot_studio(workbook: Workbook) -> None:
    """Build a selector-driven plotting workspace for routine notebook data."""

    if "Plot Studio" not in workbook.sheetnames:
        return
    worksheet = workbook["Plot Studio"]
    selected_run = worksheet["B3"].value or "EP-001"
    selected_metric = worksheet["B4"].value or "Particle size (nm)"
    selected_measurement = worksheet["B5"].value or "DLS particle size"
    for row in worksheet.iter_rows(min_row=1, max_row=max(60, worksheet.max_row), min_col=1, max_col=8):
        for cell in row:
            cell.value = None
            cell.fill = PatternFill(fill_type=None)
            cell.font = Font(name="Arial", size=10, color="1F2933")
            cell.border = Border()
    worksheet._charts = []
    worksheet.auto_filter.ref = None
    worksheet.sheet_view.showGridLines = False
    worksheet.sheet_view.selection = [Selection(activeCell="A1", sqref="A1")]
    worksheet.freeze_panes = "A7"
    worksheet.sheet_properties.tabColor = "6650A3"
    worksheet["A1"] = "PLOT STUDIO"
    worksheet["A2"] = "Choose a run and metric. Charts update automatically as Bench Log and Measurements rows are added."
    worksheet["A3"] = "Run ID"
    worksheet["B3"] = selected_run
    worksheet["A4"] = "Process metric"
    worksheet["B4"] = selected_metric
    worksheet["A5"] = "Measurement"
    worksheet["B5"] = selected_measurement
    worksheet["A7"] = "PROCESS TREND"
    worksheet["A8"] = "Timestamp"
    worksheet["B8"] = "Value"
    worksheet["D7"] = "MEASUREMENT TREND"
    worksheet["D8"] = "Sample"
    worksheet["E8"] = "Value"
    worksheet["G8"] = "Available measurements"
    metric_array = "{" + ",".join(f'\"{value}\"' for value in PROCESS_METRICS) + "}"
    process_values = "CHOOSE(MATCH($B$4," + metric_array + ",0)," + ",".join(
        f"'Bench Log'!${column}$2:${column}$1000" for column in ("D", "E", "F", "G", "H", "I", "J")
    ) + ")"
    for output_row in range(9, 101):
        process_index = excel_plot_row_index("Bench Log", process_values, output_row, "A")
        measurement_index = excel_plot_row_index(
            "Measurements", "'Measurements'!$E$2:$E$1000", output_row, "D"
        )
        choice_index = excel_plot_row_index(
            "Measurements", "'Measurements'!$C$2:$C$1000", output_row, "G"
        )
        worksheet[f"A{output_row}"] = f'=IFERROR(INDEX(\'Bench Log\'!$B$2:$B$1000,{process_index}),"")'
        worksheet[f"B{output_row}"] = f'=IFERROR(INDEX({process_values},{process_index}),"")'
        worksheet[f"D{output_row}"] = f'=IFERROR(INDEX(\'Measurements\'!$B$2:$B$1000,{measurement_index}),"")'
        worksheet[f"E{output_row}"] = f'=IFERROR(INDEX(\'Measurements\'!$E$2:$E$1000,{measurement_index}),"")'
        worksheet[f"G{output_row}"] = f'=IFERROR(INDEX(\'Measurements\'!$C$2:$C$1000,{choice_index}),"")'

    for row in (1, 7):
        for cell in worksheet[row][:8]:
            cell.fill = PatternFill("solid", fgColor=CONSOLE_NAVY if row == 1 else "6650A3")
            cell.font = Font(name="Arial", size=18 if row == 1 else 11, bold=True, color="FFFFFF")
    for cell in worksheet[2][:8]:
        cell.fill = PatternFill("solid", fgColor=CONSOLE_PALE_BLUE)
    for cell in (worksheet["B3"], worksheet["B4"], worksheet["B5"]):
        cell.fill = PatternFill("solid", fgColor=CONSOLE_INPUT)
        cell.font = Font(name="Arial", size=11, bold=True, color="1F2933")

    run_validation = DataValidation(type="list", formula1="'Experiments'!$A$2:$A$1000", allow_blank=False)
    metric_validation = DataValidation(type="list", formula1='"' + ",".join(PROCESS_METRICS) + '"', allow_blank=False)
    measurement_validation = DataValidation(type="list", formula1="'Plot Studio'!$G$9:$G$100", allow_blank=True)
    for validation, target in (
        (run_validation, "B3"),
        (metric_validation, "B4"),
        (measurement_validation, "B5"),
    ):
        worksheet.add_data_validation(validation)
        validation.add(target)

    process_chart = LineChart()
    process_chart.title = "Selected process metric"
    process_chart.y_axis.title = "Value"
    process_chart.x_axis.title = "Timestamp"
    process_chart.height = 8
    process_chart.width = 15
    process_chart.add_data(Reference(worksheet, min_col=2, min_row=8, max_row=100), titles_from_data=True)
    process_chart.set_categories(Reference(worksheet, min_col=1, min_row=9, max_row=100))
    worksheet.add_chart(process_chart, "A12")

    measurement_chart = BarChart()
    measurement_chart.type = "col"
    measurement_chart.title = "Selected measurement"
    measurement_chart.y_axis.title = "Value"
    measurement_chart.x_axis.title = "Sample"
    measurement_chart.height = 8
    measurement_chart.width = 15
    measurement_chart.add_data(Reference(worksheet, min_col=5, min_row=8, max_row=100), titles_from_data=True)
    measurement_chart.set_categories(Reference(worksheet, min_col=4, min_row=9, max_row=100))
    worksheet.add_chart(measurement_chart, "I12")

    for column, width in {"A": 22, "B": 22, "C": 3, "D": 22, "E": 18, "F": 3, "G": 26, "H": 3}.items():
        worksheet.column_dimensions[column].width = width
    worksheet.row_dimensions[1].height = 34
    worksheet.row_dimensions[2].height = 28


def build_workbook(include_examples: bool = True) -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)

    ensure_run_console(workbook)

    for spec in SHEETS:
        worksheet = workbook.create_sheet(spec.name)
        worksheet.append(list(spec.headers))
        if include_examples:
            for row in spec.example_rows:
                worksheet.append(list(row))

        for cell in worksheet[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            column = spec.columns[cell.column - 1]
            required = " Required." if column.required else ""
            cell.comment = Comment(f"{column.description}{required}", "lab-notebook-agent")

        worksheet.freeze_panes = freeze_pane_for_sheet(spec.name)
        worksheet.auto_filter.ref = worksheet.dimensions
        worksheet.sheet_view.showGridLines = True
        for column_cells in worksheet.columns:
            header = str(column_cells[0].value)
            worksheet.column_dimensions[column_cells[0].column_letter].width = (
                column_width(header)
            )
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                pattern = column_number_format(spec.name, spec.headers[cell.column - 1])
                if pattern:
                    cell.number_format = pattern

    add_validations(workbook)
    add_quality_conditional_formats(workbook)
    add_workflow_note(workbook)
    apply_workbook_presentation(workbook)
    return workbook


def save_workbook(path: str | Path, include_examples: bool = True) -> Path:
    output = Path(path).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = build_workbook(include_examples=include_examples)
    workbook.save(output)
    return output


def add_validations(workbook: Workbook) -> None:
    for sheet_name, fields in CONTROLLED_VOCAB_VALIDATIONS.items():
        worksheet = workbook[sheet_name]
        headers = [cell.value for cell in worksheet[1]]
        for field, allowed_values in fields.items():
            column_letter = get_column_letter(headers.index(field) + 1)
            formula = '"' + ",".join(allowed_values) + '"'
            validation = DataValidation(type="list", formula1=formula, allow_blank=True)
            validation.error = "Choose a value from the controlled vocabulary."
            validation.errorTitle = "Invalid value"
            validation.prompt = "Use the controlled vocabulary for this field."
            validation.promptTitle = "Controlled vocabulary"
            worksheet.add_data_validation(validation)
            validation.add(f"{column_letter}2:{column_letter}1000")

    batch_builder = workbook["Batch Builder"]
    batch_headers = [cell.value for cell in batch_builder[1]]
    for field, source_range, prompt in (
        ("experiment_id", "'Experiments'!$A$2:$A$1000", "Choose an experiment ID."),
        ("reagent_id", "'Master Reagents'!$A$2:$A$1000", "Choose a reagent ID."),
    ):
        column_letter = get_column_letter(batch_headers.index(field) + 1)
        validation = DataValidation(type="list", formula1=source_range, allow_blank=True)
        validation.error = prompt
        validation.errorTitle = "Unknown ID"
        validation.prompt = prompt
        validation.promptTitle = field.replace("_", " ").title()
        batch_builder.add_data_validation(validation)
        validation.add(f"{column_letter}2:{column_letter}1000")

    for sheet_name in ("Bench Log", "Measurements"):
        worksheet = workbook[sheet_name]
        headers = [cell.value for cell in worksheet[1]]
        run_column = get_column_letter(headers.index("Run ID") + 1)
        validation = DataValidation(
            type="list", formula1="'Experiments'!$A$2:$A$1000", allow_blank=True
        )
        validation.error = "Choose an existing run ID."
        validation.errorTitle = "Unknown run"
        worksheet.add_data_validation(validation)
        validation.add(f"{run_column}2:{run_column}1000")

    for sheet_name, field, source_range, prompt in (
        ("Experiments", "template_id", "'Experiment Templates'!$A$2:$A$1000", "Choose a governed template ID."),
        ("Inventory Transactions", "reagent_id", "'Master Reagents'!$A$2:$A$1000", "Choose an existing reagent ID."),
        ("Inventory Transactions", "experiment_id", "'Experiments'!$A$2:$A$1000", "Choose an existing experiment ID."),
        ("Equipment Bookings", "equipment_id", "'Equipment'!$A$2:$A$1000", "Choose an existing equipment ID."),
        ("Equipment Bookings", "experiment_id", "'Experiments'!$A$2:$A$1000", "Choose an existing experiment ID."),
        ("Record Signatures", "experiment_id", "'Experiments'!$A$2:$A$1000", "Choose an existing experiment ID."),
        ("Notebook Sections", "experiment_id", "'Experiments'!$A$2:$A$1000", "Choose an existing experiment ID."),
        ("Reaction Outcomes", "experiment_id", "'Experiments'!$A$2:$A$1000", "Choose an existing experiment ID."),
    ):
        worksheet = workbook[sheet_name]
        headers = [cell.value for cell in worksheet[1]]
        column_letter = get_column_letter(headers.index(field) + 1)
        validation = DataValidation(type="list", formula1=source_range, allow_blank=True)
        validation.error = prompt
        validation.errorTitle = "Unknown ID"
        validation.prompt = prompt
        worksheet.add_data_validation(validation)
        validation.add(f"{column_letter}2:{column_letter}1000")


def add_workflow_note(workbook: Workbook) -> None:
    worksheet = workbook["Agent Config"]
    worksheet.append(
        (
            "workflow_note",
            (
                "Enter reagents in Master Reagents, one experiment row in "
                "Experiments, staged quantities in Batch Builder, observations "
                "in Bench Log, required narrative in Notebook Sections, measurements "
                "in Measurements, and yield/material closure in Reaction Outcomes."
            ),
            (
                "Agent Suggestions should be treated as drafts until reviewed "
                "by a human."
            ),
        )
    )
    for cell in worksheet[worksheet.max_row]:
        cell.fill = SUBTLE_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    worksheet["A1"].comment = Comment(
        "Enter reagents in Master Reagents, one experiment row in Experiments, "
        "staged quantities in Batch Builder, observations in Bench Log, required "
        "narrative in Notebook Sections, measurements in Measurements, and yield/material "
        "closure in Reaction Outcomes. Agent Suggestions should be treated as drafts "
        "until reviewed by a human.",
        "lab-notebook-agent",
    )


def add_quality_conditional_formats(workbook: Workbook) -> None:
    results = workbook["Results"]
    results_headers = [cell.value for cell in results[1]]
    qc_column = get_column_letter(results_headers.index("qc_status") + 1)
    for status, color in (
        ("pass", "C6E0B4"),
        ("warn", "FFE699"),
        ("fail", "F4CCCC"),
        ("not_evaluated", "D9D9D9"),
    ):
        results.conditional_formatting.add(
            f"{qc_column}2:{qc_column}1000",
            CellIsRule(
                operator="equal",
                formula=[f'"{status}"'],
                fill=PatternFill("solid", fgColor=color),
            ),
        )

    deviations = workbook["Deviations"]
    deviation_headers = [cell.value for cell in deviations[1]]
    status_column = get_column_letter(deviation_headers.index("status") + 1)
    for status, color in (
        ("open", "F4CCCC"),
        ("under_review", "FFE699"),
        ("closed", "C6E0B4"),
    ):
        deviations.conditional_formatting.add(
            f"{status_column}2:{status_column}1000",
            CellIsRule(
                operator="equal",
                formula=[f'"{status}"'],
                fill=PatternFill("solid", fgColor=color),
            ),
        )

    for sheet_name, field, colors in (
        ("Experiment Templates", "state", (("effective", "C6E0B4"), ("draft", "FFE699"), ("withdrawn", "D9D9D9"))),
        ("Record Signatures", "status", (("signed", "C6E0B4"), ("draft", "FFE699"), ("revoked", "F4CCCC"))),
        ("Equipment Bookings", "status", (("confirmed", "C6E0B4"), ("planned", "FFE699"), ("cancelled", "D9D9D9"))),
        ("Reaction Master", "Governance status", (("READY_TO_ARCHIVE", "C6E0B4"), ("SIGNOFF_REQUIRED", "FFE699"), ("WITNESS_REQUIRED", "FFE699"), ("TEMPLATE_NOT_EFFECTIVE", "F4CCCC"))),
    ):
        worksheet = workbook[sheet_name]
        headers = [cell.value for cell in worksheet[1]]
        column = get_column_letter(headers.index(field) + 1)
        for status, color in colors:
            worksheet.conditional_formatting.add(
                f"{column}2:{column}1000",
                CellIsRule(operator="equal", formula=[f'"{status}"'], fill=PatternFill("solid", fgColor=color)),
            )
