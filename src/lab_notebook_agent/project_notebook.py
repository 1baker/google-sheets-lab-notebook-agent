from __future__ import annotations

import hashlib
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from .schema import SHEETS, WORKBOOK_CONTRACT_VERSION, sheet_by_name


PARSER_VERSION = "cochran-project-notebook.v2"
PROJECT_RECORD_SHEET = "Project Notebook Records"
SOURCE_SYNC_SHEET = "Source Sync"
SUPPORTED_PROFILES = (
    "auto",
    "cochran-emulsion",
    "cochran-emulsion-ledger",
    "cochran-charge-sheet",
    "cochran-charge-index",
)

_STAGE_PATTERNS = (
    ("functional shell", "functional shell"),
    ("seed", "seed"),
    ("core", "core"),
    ("shell", "shell"),
    ("workup", "workup"),
    ("chase", "chase"),
)
_ROLE_PATTERNS = (
    (
        (
            "initiator",
            "persulfate",
            "aps",
            "kps",
            "sfs",
            "ambn",
            "tert-butyl hydroperoxide",
            "tbph",
        ),
        "initiator",
    ),
    (
        (
            "surfactant",
            "aerosol",
            "sds",
            "sls",
            "dodecyl sulfate",
            "brij",
            "tween",
            "castor oil",
            "ethoxylate",
        ),
        "surfactant",
    ),
    (("buffer", "bicarbonate", "carbonate", "phosphate"), "buffer"),
    (("chain transfer", "cta", "mercaptan", "dodecyl mercaptan"), "chain_transfer_agent"),
    (("crosslink", "egdim", "divinyl", "allyl methacrylate"), "crosslinker"),
    (("water", "di water", "deionized"), "solvent"),
    (("neutraliz", "ammonia", "hydroxide"), "neutralizer"),
)
_IMPORTANT_METRICS = (
    (("particle size", "z-average", "z average"), "particle_size_nm", "nm", "result"),
    (("temperature", "temp"), "temperature_C", "C", "process_parameter"),
    (("rpm", "agitation", "stir speed"), "rpm", "rpm", "process_parameter"),
    (("feed rate",), "feed_rate_mL_min", "mL/min", "process_parameter"),
    (("radical flux",), "radical_flux_mol_min_L", "mol/min/L", "process_parameter"),
    (("solids",), "", "%", "result"),
    (("conversion",), "", "%", "result"),
    (("particle count", "particles"), "", "", "process_parameter"),
    (("seed mass", "water mass", "total mass", "total volume"), "", "", "process_parameter"),
    (("surfactant", "cmc"), "", "", "process_parameter"),
    (("moles", "ratio", "diameter", "hold time"), "", "", "process_parameter"),
    (("ph",), "", "", "result"),
)
_NOTE_TERMS = (
    "observ",
    "looked",
    "appeared",
    "separat",
    "coag",
    "foam",
    "spray",
    "shake",
    "freeze",
    "pass through",
    "high shear",
    "heat",
    "hold",
    "wait",
    "add ",
    "mixed",
    "peak",
    "equilibr",
)


def parse_project_notebook_sheet(
    values: list[list[Any]],
    *,
    source_spreadsheet_id: str,
    source_title: str,
    source_sheet_name: str,
    source_sheet_id: int | str | None = None,
    source_url: str = "",
    source_modified_at: str = "",
    source_range: str = "A1:AO1200",
    parser_profile: str = "auto",
    synced_at: str | None = None,
) -> dict[str, Any]:
    if parser_profile not in SUPPORTED_PROFILES:
        raise ValueError(f"Unsupported parser profile {parser_profile!r}.")
    normalized_values = normalize_values(values)
    selected_profile = (
        detect_profile(normalized_values)
        if parser_profile == "auto"
        else parser_profile
    )
    source_fingerprint = fingerprint_values(normalized_values)
    source_key = build_source_key(
        source_spreadsheet_id,
        source_sheet_name,
        source_sheet_id,
    )
    synced_at = synced_at or utc_now()
    source_context = {
        "source_key": source_key,
        "source_fingerprint": source_fingerprint,
        "source_range": source_range,
        "synced_at": synced_at,
    }
    experiment_id = experiment_id_from_sheet_name(source_sheet_name)
    lineage_experiments: list[dict[str, str]] = []
    if selected_profile == "cochran-emulsion-ledger":
        records, warnings = parse_emulsion_ledger_records(
            normalized_values,
            experiment_id,
            source_context,
        )
    elif selected_profile == "cochran-charge-sheet":
        records, warnings = parse_charge_sheet_records(
            normalized_values,
            experiment_id,
            source_context,
        )
    elif selected_profile == "cochran-charge-index":
        records, warnings, lineage_experiments = parse_charge_index_records(
            normalized_values,
            source_context,
        )
    else:
        records, warnings = parse_emulsion_recipe_records(
            normalized_values,
            experiment_id,
            source_context,
        )

    source_reference = source_url or (
        f"https://docs.google.com/spreadsheets/d/{source_spreadsheet_id}"
        if source_spreadsheet_id
        else ""
    )
    experiment = project_row(
        "Experiments",
        {
            "experiment_id": experiment_id,
            "date": find_source_date(normalized_values),
            "project": "Cochran Research Group",
            "process_type": "emulsion polymerization",
            "objective": f"Import and review the recipe, process, feed schedule, and results from {source_sheet_name}.",
            "status": "needs_review",
            "summary": (
                f"Imported {len(records)} normalized records from {source_title or source_spreadsheet_id}, "
                f"tab {source_sheet_name}."
            ),
            "source_notebook_id": source_spreadsheet_id,
            "source_sheet_name": source_sheet_name,
            "source_url": source_reference,
            "source_modified_at": source_modified_at,
            "source_fingerprint": source_fingerprint,
        },
    )
    experiments = (
        [
            project_row(
                "Experiments",
                {
                    "experiment_id": item["experiment_id"],
                    "date": item.get("date", ""),
                    "project": "Cochran Research Group",
                    "process_type": "emulsion polymerization",
                    "objective": (
                        "Review material-lot lineage and product provenance for "
                        f"{item.get('description', item['experiment_id'])}."
                    ),
                    "status": "needs_review",
                    "summary": (
                        f"Imported lineage from {source_title or source_spreadsheet_id}, "
                        f"tab {source_sheet_name}."
                    ),
                    "source_notebook_id": source_spreadsheet_id,
                    "source_sheet_name": source_sheet_name,
                    "source_url": source_reference,
                    "source_modified_at": source_modified_at,
                    "source_fingerprint": source_fingerprint,
                },
            )
            for item in lineage_experiments
        ]
        if lineage_experiments
        else [experiment]
    )
    sync_row = project_row(
        SOURCE_SYNC_SHEET,
        {
            "source_key": source_key,
            "source_spreadsheet_id": source_spreadsheet_id,
            "source_title": source_title,
            "source_sheet_name": source_sheet_name,
            "source_sheet_id": source_sheet_id if source_sheet_id is not None else "",
            "source_url": source_reference,
            "source_modified_at": source_modified_at,
            "source_range": source_range,
            "parser_profile": selected_profile,
            "parser_version": PARSER_VERSION,
            "source_fingerprint": source_fingerprint,
            "synced_at": synced_at,
            "record_count": len(records),
            "status": "synced" if records else "needs_review",
            "warnings_json": json.dumps(warnings, sort_keys=True),
        },
    )
    return {
        "source_key": source_key,
        "parser_profile": selected_profile,
        "source_fingerprint": source_fingerprint,
        "experiment": experiments[0],
        "experiments": experiments,
        "records": records,
        "source_sync": sync_row,
        "warnings": warnings,
    }


def parse_project_notebook_workbook(
    workbook_path: str | Path,
    source_sheet_names: tuple[str, ...],
    *,
    source_range: str = "A1:AO1200",
    parser_profile: str = "auto",
    synced_at: str | None = None,
) -> list[dict[str, Any]]:
    if not source_sheet_names:
        raise ValueError("At least one source sheet is required.")
    if len(set(source_sheet_names)) != len(source_sheet_names):
        raise ValueError("Source sheet names must be unique.")
    source = Path(workbook_path).expanduser().resolve()
    workbook = load_workbook(source, read_only=True, data_only=True)
    missing = [name for name in source_sheet_names if name not in workbook.sheetnames]
    if missing:
        workbook.close()
        raise ValueError("Source tabs were not found: " + ", ".join(missing))
    min_column, min_row, max_column, max_row = range_boundaries(source_range)
    modified_at = datetime.fromtimestamp(
        source.stat().st_mtime,
        timezone.utc,
    ).replace(microsecond=0).isoformat()
    parsed_sources = []
    try:
        for sheet_name in source_sheet_names:
            worksheet = workbook[sheet_name]
            if worksheet.max_row < min_row or worksheet.max_column < min_column:
                values = []
            else:
                values = [
                    list(row)
                    for row in worksheet.iter_rows(
                        min_row=min_row,
                        max_row=min(max_row, worksheet.max_row),
                        min_col=min_column,
                        max_col=min(max_column, worksheet.max_column),
                        values_only=True,
                    )
                ]
            parsed_sources.append(
                parse_project_notebook_sheet(
                    values,
                    source_spreadsheet_id=str(source),
                    source_title=source.name,
                    source_sheet_name=sheet_name,
                    source_url=str(source),
                    source_modified_at=modified_at,
                    source_range=source_range,
                    parser_profile=parser_profile,
                    synced_at=synced_at,
                )
            )
    finally:
        workbook.close()
    return parsed_sources


def parse_emulsion_recipe_records(
    values: list[list[Any]],
    experiment_id: str,
    source: dict[str, str],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    records: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    current_stage = ""
    current_section = ""
    component_columns: dict[str, int] = {}
    component_header_row = -1
    feed_rows: set[int] = set()
    occurrence: dict[tuple[str, str, str, str], int] = {}

    for row_index, row in enumerate(values):
        row_text = " | ".join(cell_text(cell) for cell in row if cell_text(cell))
        stage = stage_from_row(row)
        if stage:
            current_stage = stage
            current_section = ""

        detected_columns = component_header_columns(row)
        if detected_columns:
            component_columns = detected_columns
            component_header_row = row_index
            continue

        if is_feed_header(row):
            parsed_feed_rows, consumed = parse_feed_table(
                values,
                row_index,
                experiment_id,
                current_stage,
                current_section,
                source,
                occurrence,
            )
            records.extend(parsed_feed_rows)
            feed_rows.update(consumed)
            continue

        if row_index in feed_rows:
            continue

        if component_columns and row_index > component_header_row:
            material = value_at(row, component_columns.get("material"))
            section_value = value_at(row, component_columns.get("section"))
            pphm = numeric_value(value_at(row, component_columns.get("pphm")))
            mass = numeric_value(value_at(row, component_columns.get("mass")))
            density = numeric_value(value_at(row, component_columns.get("density")))
            volume = numeric_value(value_at(row, component_columns.get("volume")))
            actual_mass = numeric_value(value_at(row, component_columns.get("actual_mass")))
            if section_value and not is_number(section_value) and not material:
                current_section = cell_text(section_value)
            elif (
                material
                and not is_number(material)
                and any(value is not None for value in (pphm, mass, volume, actual_mass))
            ):
                if section_value and not is_number(section_value):
                    current_section = cell_text(section_value)
                label = cell_text(material)
                role = infer_material_role(label, current_stage, current_section)
                planned = mass if mass is not None else (volume if volume is not None else pphm)
                units = "g" if mass is not None else ("mL" if volume is not None else "pphm")
                records.append(
                    make_record(
                        experiment_id=experiment_id,
                        record_type="component",
                        stage=current_stage,
                        section=current_section,
                        label=label,
                        source=source,
                        occurrence=occurrence,
                        material_name=label,
                        target_role=role,
                        planned_value=display_number(planned),
                        actual_value=display_number(actual_mass),
                        units=units,
                        parts_per_hundred_monomer=display_number(pphm),
                        mass_g=display_number(mass),
                        actual_mass_g=display_number(actual_mass),
                        volume_mL=display_number(volume),
                        density_g_mL=display_number(density),
                        source_range=row_source_range(source["source_range"], row_index),
                        details_json=details_json_for_row(row),
                    )
                )
                continue

        if row_index in feed_rows or not row_text:
            continue
        records.extend(
            metric_records_from_row(
                row,
                row_index,
                experiment_id,
                current_stage,
                current_section,
                source,
                occurrence,
            )
        )
        note = note_from_row(row)
        if note:
            records.append(
                make_record(
                    experiment_id=experiment_id,
                    record_type="observation",
                    stage=current_stage,
                    section=current_section,
                    label=note[:120],
                    source=source,
                    occurrence=occurrence,
                    notes=note,
                    source_range=row_source_range(source["source_range"], row_index),
                    details_json=details_json_for_row(row),
                )
            )

    if not any(record["record_type"] == "component" for record in records):
        warnings.append(
            {
                "code": "no_component_table_detected",
                "message": "No pphm/mass component table was detected; review this tab or choose the ledger profile.",
            }
        )
    return deduplicate_records(records), warnings


def parse_emulsion_ledger_records(
    values: list[list[Any]],
    experiment_id: str,
    source: dict[str, str],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    records: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    occurrence: dict[tuple[str, str, str, str], int] = {}
    header_count = 0
    for row_index, row in enumerate(values):
        columns = ledger_header_columns(row)
        if not columns:
            continue
        header_count += 1
        section = nearest_section_label(values, row_index)
        for component_index in range(row_index + 1, len(values)):
            component_row = values[component_index]
            material = cell_text(value_at(component_row, columns.get("material")))
            if not material:
                if component_index > row_index + 1:
                    break
                continue
            if normalized_label(material) in {"total", "totals", "component"}:
                break
            planned = numeric_value(value_at(component_row, columns.get("planned")))
            actual = numeric_value(value_at(component_row, columns.get("actual")))
            wt_percent = numeric_value(value_at(component_row, columns.get("wt_percent")))
            if all(value is None for value in (planned, actual, wt_percent)):
                if note_from_row(component_row):
                    break
                continue
            records.append(
                make_record(
                    experiment_id=experiment_id,
                    record_type="component",
                    stage="formulation",
                    section=section,
                    label=material,
                    source=source,
                    occurrence=occurrence,
                    material_name=material,
                    target_role=infer_material_role(material, "formulation", section),
                    planned_value=display_number(planned),
                    actual_value=display_number(actual),
                    units="g" if planned is not None or actual is not None else "%",
                    mass_g=display_number(planned),
                    actual_mass_g=display_number(actual),
                    concentration=display_number(wt_percent),
                    concentration_units="wt fraction" if wt_percent is not None and wt_percent <= 1 else "wt%",
                    source_range=row_source_range(source["source_range"], component_index),
                    details_json=details_json_for_row(component_row),
                )
            )

    for row_index, row in enumerate(values):
        note = note_from_row(row)
        if not note:
            continue
        records.append(
            make_record(
                experiment_id=experiment_id,
                record_type="observation",
                stage="process",
                section="",
                label=note[:120],
                source=source,
                occurrence=occurrence,
                notes=note,
                source_range=row_source_range(source["source_range"], row_index),
                details_json=details_json_for_row(row),
            )
        )
    if not header_count:
        warnings.append(
            {
                "code": "no_ledger_table_detected",
                "message": "No component/goal/actual table was detected; review this tab or choose the recipe profile.",
            }
        )
    return deduplicate_records(records), warnings


def parse_charge_sheet_records(
    values: list[list[Any]],
    experiment_id: str,
    source: dict[str, str],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    records: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    occurrence: dict[tuple[str, str, str, str], int] = {}
    parameter_section = ""
    charge_stage = ""
    process_log_columns: dict[str, int] = {}
    process_log_torque_units = ""
    process_log_rows = 0

    for row_index, row in enumerate(values):
        labels = [normalized_label(value) for value in row]
        step_cell = next(
            (
                cell_text(value)
                for value in row
                if re.search(r"\bstep\s+\d+", cell_text(value), re.IGNORECASE)
            ),
            "",
        )
        if step_cell:
            match = re.search(r"\bstep\s+(\d+)", step_cell, re.IGNORECASE)
            charge_stage = f"step {match.group(1)}" if match else charge_stage

        detected_log_columns = charge_process_log_columns(row)
        if detected_log_columns:
            process_log_columns = detected_log_columns
            process_log_torque_units = header_units(
                value_at(row, process_log_columns.get("torque"))
            )
            continue
        if process_log_columns:
            source_time = cell_text(value_at(row, process_log_columns.get("time")))
            temperature = numeric_value(value_at(row, process_log_columns.get("temperature")))
            rpm = numeric_value(value_at(row, process_log_columns.get("rpm")))
            oil_temperature = numeric_value(
                value_at(row, process_log_columns.get("oil_temperature"))
            )
            torque = numeric_value(value_at(row, process_log_columns.get("torque")))
            raw_addition = value_at(row, process_log_columns.get("addition"))
            addition = numeric_value(raw_addition)
            raw_cumulative_addition = value_at(
                row, process_log_columns.get("cumulative_addition")
            )
            cumulative_addition = numeric_value(raw_cumulative_addition)
            note = cell_text(value_at(row, process_log_columns.get("note")))
            if not note and addition is None:
                note = cell_text(raw_addition)
            if not note and cumulative_addition is None:
                note = cell_text(raw_cumulative_addition)
            if source_time and any(
                value is not None
                for value in (
                    temperature,
                    rpm,
                    oil_temperature,
                    torque,
                    addition,
                    cumulative_addition,
                )
            ) or (source_time and note):
                duration = duration_minutes(
                    value_at(row, process_log_columns.get("duration"))
                )
                records.append(
                    make_record(
                        experiment_id=experiment_id,
                        record_type="observation",
                        stage="polymerization",
                        section="process log",
                        label=source_time,
                        source=source,
                        occurrence=occurrence,
                        actual_value=display_number(addition),
                        units="g" if addition is not None else "",
                        elapsed_end_min=display_number(duration),
                        temperature_C=display_number(temperature),
                        rpm=display_number(rpm),
                        oil_temperature_C=display_number(oil_temperature),
                        torque=display_number(torque),
                        torque_units=process_log_torque_units,
                        cumulative_addition_g=display_number(cumulative_addition),
                        source_timestamp=source_time,
                        notes=note,
                        source_range=row_source_range(source["source_range"], row_index),
                        details_json=details_json_for_row(row),
                    )
                )
                process_log_rows += 1
                continue

        material = cell_text(value_at(row, 5))
        planned_mass = numeric_value(value_at(row, 6))
        actual_mass = numeric_value(value_at(row, 7))
        mass_variance = numeric_value(value_at(row, 8))
        units = cell_text(value_at(row, 9))
        if (
            material
            and normalized_label(material) not in {"total"}
            and any(
                value is not None
                for value in (planned_mass, actual_mass, mass_variance)
            )
            and "step " not in normalized_label(material)
        ):
            records.append(
                make_record(
                    experiment_id=experiment_id,
                    record_type="component",
                    stage=charge_stage,
                    section="mass charge",
                    label=material,
                    source=source,
                    occurrence=occurrence,
                    material_name=material,
                    target_role=infer_material_role(material, charge_stage, "mass charge"),
                    planned_value=display_number(planned_mass),
                    actual_value=display_number(actual_mass),
                    units=units or "g",
                    mass_g=display_number(planned_mass),
                    actual_mass_g=display_number(actual_mass),
                    mass_variance_g=display_number(mass_variance),
                    source_range=row_source_range(source["source_range"], row_index),
                    details_json=details_json_for_row(row),
                )
            )

        parameter_label = cell_text(value_at(row, 1))
        parameter_value = value_at(row, 2)
        parameter_units = cell_text(value_at(row, 3))
        if parameter_label and not cell_text(parameter_value):
            if any(
                term in normalized_label(parameter_label)
                for term in ("parameter", "composition", "scale")
            ):
                parameter_section = parameter_label
        elif parameter_label and cell_text(parameter_value):
            records.append(
                make_record(
                    experiment_id=experiment_id,
                    record_type="process_parameter",
                    stage="setup",
                    section=parameter_section,
                    label=parameter_label,
                    source=source,
                    occurrence=occurrence,
                    planned_value=cell_text(parameter_value),
                    units=parameter_units,
                    source_range=row_source_range(source["source_range"], row_index),
                    details_json=details_json_for_row(row),
                )
            )

        for value in row:
            instruction = cell_text(value)
            if len(instruction) < 32:
                continue
            normalized = normalized_label(instruction)
            if (
                "\n" in instruction
                or any(
                    term in normalized
                    for term in (
                        "combine",
                        "argon",
                        "inject",
                        "terminate",
                        "polymerization",
                        "temperature",
                    )
                )
            ):
                records.append(
                    make_record(
                        experiment_id=experiment_id,
                        record_type="observation",
                        stage=charge_stage or "procedure",
                        section="instruction",
                        label=instruction.splitlines()[0][:120],
                        source=source,
                        occurrence=occurrence,
                        notes=instruction,
                        source_range=row_source_range(source["source_range"], row_index),
                        details_json=details_json_for_row(row),
                    )
                )
    if not any(record["record_type"] == "component" for record in records):
        warnings.append(
            {
                "code": "no_charge_table_detected",
                "message": "No planned/actual mass-charge rows were detected.",
            }
        )
    if not process_log_rows:
        warnings.append(
            {
                "code": "no_process_log_detected",
                "message": "No time/temperature/RPM process log was detected.",
            }
        )
    return deduplicate_records(records), warnings


def parse_charge_index_records(
    values: list[list[Any]],
    source: dict[str, str],
) -> tuple[
    list[dict[str, Any]],
    list[dict[str, Any]],
    list[dict[str, str]],
]:
    records: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    experiments: list[dict[str, str]] = []
    experiment_ids: set[str] = set()
    occurrence: dict[tuple[str, str, str, str], int] = {}
    lot_headers: dict[int, str] = {}
    data_start = -1

    for row_index, row in enumerate(values):
        if normalized_label(value_at(row, 0)).startswith(
            "experiment name/description"
        ):
            header_row = values[row_index + 1] if row_index + 1 < len(values) else []
            lot_headers = {
                index: cell_text(value)
                for index, value in enumerate(header_row)
                if index >= 2 and cell_text(value)
            }
            data_start = row_index + 2
            continue
        if data_start < 0 or row_index < data_start or not lot_headers:
            continue
        description = cell_text(value_at(row, 0))
        if not description:
            continue
        if normalized_label(description) in {"lab scale"}:
            continue
        date = normalize_date_text(value_at(row, 1))
        product_index = next(
            (
                index
                for index, header in lot_headers.items()
                if normalized_label(header) == "product"
            ),
            None,
        )
        product_lot = cell_text(value_at(row, product_index))
        experiment_id = lineage_experiment_id(description, product_lot)
        if experiment_id not in experiment_ids:
            experiment_ids.add(experiment_id)
            experiments.append(
                {
                    "experiment_id": experiment_id,
                    "date": date,
                    "description": description,
                }
            )
        for index, header in lot_headers.items():
            lot = cell_text(value_at(row, index))
            if not lot or normalized_label(lot) in {"n/a", "na"}:
                continue
            if normalized_label(header) == "notes":
                records.append(
                    make_record(
                        experiment_id=experiment_id,
                        record_type="observation",
                        stage="lineage",
                        section=description,
                        label="Master sheet note",
                        source=source,
                        occurrence=occurrence,
                        notes=lot,
                        source_range=row_source_range(source["source_range"], row_index),
                        details_json=details_json_for_row(row),
                    )
                )
            elif normalized_label(header) == "product":
                records.append(
                    make_record(
                        experiment_id=experiment_id,
                        record_type="result",
                        stage="lineage",
                        section=description,
                        label="Product lot",
                        source=source,
                        occurrence=occurrence,
                        actual_value=lot,
                        units="lot",
                        product_lot=lot,
                        source_range=row_source_range(source["source_range"], row_index),
                        details_json=details_json_for_row(row),
                    )
                )
            else:
                records.append(
                    make_record(
                        experiment_id=experiment_id,
                        record_type="component",
                        stage="lineage",
                        section=description,
                        label=header,
                        source=source,
                        occurrence=occurrence,
                        material_name=header,
                        target_role=infer_material_role(
                            header,
                            "lineage",
                            "material lot",
                        ),
                        lot_number=lot,
                        source_range=row_source_range(source["source_range"], row_index),
                        details_json=details_json_for_row(row),
                    )
                )
    if not experiments:
        warnings.append(
            {
                "code": "no_lineage_rows_detected",
                "message": "No experiment/date/material-lot rows were detected.",
            }
        )
    return deduplicate_records(records), warnings, experiments


def build_project_notebook_sync_report(
    parsed_sources: list[dict[str, Any]],
    tables: dict[str, list[dict[str, Any]]],
) -> dict[str, Any]:
    existing_experiments = indexed_rows(tables.get("Experiments", []), "experiment_id")
    existing_records = indexed_rows(tables.get(PROJECT_RECORD_SHEET, []), "record_id")
    existing_sync = indexed_rows(tables.get(SOURCE_SYNC_SHEET, []), "source_key")
    append_experiments: list[dict[str, Any]] = []
    update_experiments: list[dict[str, Any]] = []
    append_records: list[dict[str, Any]] = []
    update_records: list[dict[str, Any]] = []
    append_sync: list[dict[str, Any]] = []
    update_sync: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    unchanged_sources: list[str] = []
    changed_sources: list[str] = []
    seen_source_keys: set[str] = set()
    seen_experiment_ids: set[str] = set()

    for parsed in parsed_sources:
        source_key = str(parsed.get("source_key", "")).strip()
        if not source_key:
            raise ValueError("Every parsed source must include source_key.")
        if source_key in seen_source_keys:
            raise ValueError(f"Duplicate parsed source_key {source_key!r}.")
        seen_source_keys.add(source_key)
        desired_sync = dict(parsed.get("source_sync", {}))
        current_sync_entry = existing_sync.get(source_key)
        current_sync = current_sync_entry[1] if current_sync_entry else {}
        if (
            current_sync
            and same_cell(current_sync.get("source_fingerprint"), desired_sync.get("source_fingerprint"))
            and same_cell(current_sync.get("parser_profile"), desired_sync.get("parser_profile"))
            and same_cell(current_sync.get("parser_version"), desired_sync.get("parser_version"))
            and same_cell(current_sync.get("source_sheet_name"), desired_sync.get("source_sheet_name"))
            and same_cell(current_sync.get("source_range"), desired_sync.get("source_range"))
            and same_cell(current_sync.get("source_url"), desired_sync.get("source_url"))
            and same_cell(current_sync.get("source_modified_at"), desired_sync.get("source_modified_at"))
        ):
            unchanged_sources.append(source_key)
            continue
        changed_sources.append(source_key)
        warnings.extend(parsed.get("warnings", []))

        parsed_experiments = parsed.get("experiments")
        if not isinstance(parsed_experiments, list):
            parsed_experiments = [parsed.get("experiment", {})]
        for experiment_value in parsed_experiments:
            if not isinstance(experiment_value, dict):
                continue
            experiment = dict(experiment_value)
            experiment_id = str(experiment.get("experiment_id", "")).strip()
            if not experiment_id or experiment_id in seen_experiment_ids:
                continue
            seen_experiment_ids.add(experiment_id)
            current_experiment_entry = existing_experiments.get(experiment_id)
            if current_experiment_entry is None:
                append_experiments.append(experiment)
                continue
            row_number, current_experiment = current_experiment_entry
            source_fields = (
                "source_notebook_id",
                "source_sheet_name",
                "source_url",
                "source_modified_at",
                "source_fingerprint",
            )
            for field in source_fields:
                value = experiment.get(field, "")
                if not same_cell(current_experiment.get(field), value):
                    update_experiments.append(
                        cell_update(
                            "Experiments",
                            row_number,
                            "experiment_id",
                            experiment_id,
                            field,
                            value,
                        )
                    )
            for field in ("project", "process_type", "objective", "summary"):
                if not cell_text(current_experiment.get(field)) and cell_text(
                    experiment.get(field)
                ):
                    update_experiments.append(
                        cell_update(
                            "Experiments",
                            row_number,
                            "experiment_id",
                            experiment_id,
                            field,
                            experiment.get(field, ""),
                        )
                    )

        desired_record_ids: set[str] = set()
        for record in parsed.get("records", []) or []:
            if not isinstance(record, dict):
                continue
            record_id = str(record.get("record_id", "")).strip()
            if not record_id:
                continue
            desired_record_ids.add(record_id)
            existing_entry = existing_records.get(record_id)
            if existing_entry is None:
                append_records.append(record)
                continue
            row_number, current_record = existing_entry
            update_records.extend(
                diff_row(
                    PROJECT_RECORD_SHEET,
                    row_number,
                    "record_id",
                    record_id,
                    current_record,
                    record,
                )
            )
        for record_id, (row_number, current_record) in existing_records.items():
            if str(current_record.get("source_key", "")).strip() != source_key:
                continue
            if record_id in desired_record_ids:
                continue
            if normalized_label(current_record.get("active", "")) != "false":
                update_records.append(
                    cell_update(
                        PROJECT_RECORD_SHEET,
                        row_number,
                        "record_id",
                        record_id,
                        "active",
                        "false",
                    )
                )

        if current_sync_entry is None:
            append_sync.append(desired_sync)
        else:
            row_number, current_sync = current_sync_entry
            update_sync.extend(
                diff_row(
                    SOURCE_SYNC_SHEET,
                    row_number,
                    "source_key",
                    source_key,
                    current_sync,
                    desired_sync,
                )
            )

    report = {
        "schema": "lab-notebook-agent-project-notebook-sync.v1",
        "append_experiments": append_experiments,
        "update_experiments": update_experiments,
        "append_project_notebook_records": append_records,
        "update_project_notebook_records": update_records,
        "append_source_sync": append_sync,
        "update_source_sync": update_sync,
        "unchanged_sources": unchanged_sources,
        "changed_sources": changed_sources,
        "warnings": warnings,
    }
    report["summary"] = {
        "source_count": len(parsed_sources),
        "changed_source_count": len(changed_sources),
        "unchanged_source_count": len(unchanged_sources),
        "experiment_rows_to_append": len(append_experiments),
        "experiment_cells_to_update": len(update_experiments),
        "project_records_to_append": len(append_records),
        "project_record_cells_to_update": len(update_records),
        "source_sync_rows_to_append": len(append_sync),
        "source_sync_cells_to_update": len(update_sync),
        "warning_count": len(warnings),
    }
    return report


def apply_project_notebook_sync_report_to_workbook(
    workbook_path: str | Path,
    report: dict[str, Any],
    output_workbook: str | Path | None = None,
) -> Path:
    source = Path(workbook_path).expanduser().resolve()
    destination = Path(output_workbook).expanduser().resolve() if output_workbook else source
    workbook = load_workbook(source)
    ensure_workbook_contract(workbook)
    for sheet_name, report_key in (
        ("Experiments", "append_experiments"),
        (PROJECT_RECORD_SHEET, "append_project_notebook_records"),
        (SOURCE_SYNC_SHEET, "append_source_sync"),
    ):
        worksheet = workbook[sheet_name]
        headers = [str(cell.value or "") for cell in worksheet[1]]
        for row in report.get(report_key, []) or []:
            if isinstance(row, dict):
                worksheet.append([row.get(header, "") for header in headers])
    for sheet_name, report_key in (
        ("Experiments", "update_experiments"),
        (PROJECT_RECORD_SHEET, "update_project_notebook_records"),
        (SOURCE_SYNC_SHEET, "update_source_sync"),
    ):
        worksheet = workbook[sheet_name]
        headers = [str(cell.value or "") for cell in worksheet[1]]
        header_index = {header: index + 1 for index, header in enumerate(headers)}
        for update in report.get(report_key, []) or []:
            if not isinstance(update, dict):
                continue
            row_number = int(update.get("row_number", 0) or 0)
            field = str(update.get("field", ""))
            if row_number >= 2 and field in header_index:
                worksheet.cell(
                    row=row_number,
                    column=header_index[field],
                    value=update.get("value", ""),
                )
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    return destination


def ensure_workbook_contract(workbook: Any) -> None:
    for spec in SHEETS:
        if spec.name not in workbook.sheetnames:
            worksheet = workbook.create_sheet(spec.name)
            worksheet.append(list(spec.headers))
            worksheet.freeze_panes = "A2"
            continue
        worksheet = workbook[spec.name]
        headers = [str(cell.value or "") for cell in worksheet[1]]
        for header in spec.headers:
            if header not in headers:
                headers.append(header)
                worksheet.cell(row=1, column=len(headers), value=header)
    seed_local_contract_reference_rows(workbook)


def seed_local_contract_reference_rows(workbook: Any) -> None:
    keyed_seed_specs = {
        "Process Knowledge": ("process_type", "material_role"),
        "Controlled Vocab": ("field", "allowed_value"),
        "Agent Config": ("key",),
    }
    for sheet_name, key_fields in keyed_seed_specs.items():
        worksheet = workbook[sheet_name]
        headers = [str(cell.value or "") for cell in worksheet[1]]
        existing_keys = {
            tuple(str(worksheet.cell(row=row, column=headers.index(field) + 1).value or "") for field in key_fields)
            for row in range(2, worksheet.max_row + 1)
        }
        spec = sheet_by_name(sheet_name)
        for values in spec.example_rows:
            row = {
                header: values[index] if index < len(values) else ""
                for index, header in enumerate(spec.headers)
            }
            key = tuple(str(row.get(field, "")) for field in key_fields)
            if key in existing_keys:
                continue
            worksheet.append([row.get(header, "") for header in spec.headers])
            existing_keys.add(key)

    metadata = workbook["Workbook Metadata"]
    headers = [str(cell.value or "") for cell in metadata[1]]
    rows_by_key = {
        str(metadata.cell(row=row, column=headers.index("key") + 1).value or ""): row
        for row in range(2, metadata.max_row + 1)
    }
    migrated_at = datetime.now(timezone.utc).isoformat()
    desired = {
        "contract_name": "lab-notebook-agent-workbook",
        "contract_version": WORKBOOK_CONTRACT_VERSION,
        "workbook_timezone": "America/Chicago",
        "migration_status": "current",
        "last_migrated_at": migrated_at,
    }
    for key, value in desired.items():
        row_number = rows_by_key.get(key)
        if row_number is None:
            metadata.append([key, value, migrated_at, "Managed workbook contract metadata."])
            continue
        metadata.cell(row=row_number, column=headers.index("value") + 1, value=value)
        metadata.cell(row=row_number, column=headers.index("updated_at") + 1, value=migrated_at)


def detect_profile(values: list[list[Any]]) -> str:
    flattened = " ".join(
        normalized_label(value)
        for row in values[:80]
        for value in row
        if cell_text(value)
    )
    if (
        "experiment name/description" in flattened
        and "lot number" in flattened
        and "product" in flattened
    ):
        return "cochran-charge-index"
    if (
        "mass charge values" in flattened
        and "rxn temp" in flattened
        and "actual" in flattened
    ):
        return "cochran-charge-sheet"
    for row in values:
        labels = {normalized_label(cell) for cell in row if cell_text(cell)}
        if (
            any(label in {"component", "material"} for label in labels)
            and any("actual" in label for label in labels)
            and any("goal" in label or "target" in label for label in labels)
        ):
            return "cochran-emulsion-ledger"
    return "cochran-emulsion"


def charge_process_log_columns(row: list[Any]) -> dict[str, int]:
    labels = [normalized_label(value) for value in row]
    time_index = first_index(labels, lambda label: label == "time")
    temperature_index = first_index(
        labels,
        lambda label: "rxn temp" in label or "reaction temp" in label,
    )
    rpm_index = first_index(labels, lambda label: label == "rpm")
    if time_index is None or temperature_index is None or rpm_index is None:
        return {}
    columns = {
        "time": time_index,
        "temperature": temperature_index,
        "rpm": rpm_index,
    }
    for name, matcher in (
        ("duration", lambda label: label == "duration"),
        ("oil_temperature", lambda label: "oil temp" in label),
        ("torque", lambda label: "torque" in label),
        ("addition", lambda label: "initiator" in label and "added" in label),
        (
            "cumulative_addition",
            lambda label: "initiator" in label
            and ("total" in label or "cumulative" in label),
        ),
    ):
        index = first_index(labels, matcher)
        if index is not None:
            columns[name] = index
    torque_index = columns.get("torque")
    if torque_index is not None:
        if "addition" not in columns and torque_index + 1 < len(row):
            columns["addition"] = torque_index + 1
        if "cumulative_addition" not in columns and torque_index + 2 < len(row):
            columns["cumulative_addition"] = torque_index + 2
    if len(row) > max(columns.values()) + 1:
        columns["note"] = max(columns.values()) + 1
    return columns


def header_units(value: Any) -> str:
    text = cell_text(value)
    match = re.search(r"\(([^()]+)\)", text)
    return match.group(1).strip() if match else ""


def component_header_columns(row: list[Any]) -> dict[str, int]:
    labels = [normalized_label(value) for value in row]
    pphm_index = first_index(labels, lambda label: "parts per hundred monomer" in label or label == "pphm")
    if pphm_index is None:
        return {}
    material_index = pphm_index - 1
    if material_index < 0:
        return {}
    columns = {"material": material_index, "pphm": pphm_index}
    if material_index > 0:
        columns["section"] = material_index - 1
    for name, matcher in (
        ("actual_mass", lambda label: "actual" in label and ("mass" in label or label == "actual")),
        ("mass", lambda label: ("mass" in label or "goal" in label) and "actual" not in label),
        ("density", lambda label: "density" in label),
        ("volume", lambda label: "volume" in label),
    ):
        index = first_index(labels, matcher)
        if index is not None:
            columns[name] = index
    return columns


def ledger_header_columns(row: list[Any]) -> dict[str, int]:
    labels = [normalized_label(value) for value in row]
    material = first_index(labels, lambda label: label in {"component", "material", "ingredients", "ingredient"})
    actual = first_index(labels, lambda label: "actual" in label)
    planned = first_index(labels, lambda label: "goal" in label or "target" in label or "planned" in label)
    if material is None or actual is None or planned is None:
        return {}
    columns = {"material": material, "actual": actual, "planned": planned}
    wt_percent = first_index(labels, lambda label: "wt" in label and ("%" in label or "percent" in label))
    if wt_percent is not None:
        columns["wt_percent"] = wt_percent
    return columns


def parse_feed_table(
    values: list[list[Any]],
    header_row_index: int,
    experiment_id: str,
    stage: str,
    section: str,
    source: dict[str, str],
    occurrence: dict[tuple[str, str, str, str], int],
) -> tuple[list[dict[str, Any]], set[int]]:
    header = values[header_row_index]
    header_labels = [cell_text(value) for value in header]
    records: list[dict[str, Any]] = []
    consumed = {header_row_index}
    blank_count = 0
    for row_index in range(header_row_index + 1, len(values)):
        row = values[row_index]
        if stage_from_row(row) or component_header_columns(row) or is_feed_header(row):
            break
        if not any(cell_text(value) for value in row):
            blank_count += 1
            if blank_count >= 2:
                break
            continue
        blank_count = 0
        interval_index = first_index(
            [normalized_label(label) for label in header_labels],
            lambda label: "time interval" in label,
        )
        interval = cell_text(value_at(row, interval_index))
        if not interval or not re.search(r"\d", interval):
            continue
        start_min, end_min = parse_time_interval(interval)
        details = {
            header_labels[index] or f"column_{index + 1}": value
            for index, value in enumerate(row)
            if cell_text(value)
        }
        feed_rate = first_numeric_by_header(details, ("feed rate",), exclude=("oxidant", "reductant"))
        cumulative = first_numeric_by_header(details, ("cumulative", "%"))
        radical_flux = first_numeric_by_header(details, ("radical flux",))
        particle_size = first_numeric_by_header(details, ("particle size",))
        records.append(
            make_record(
                experiment_id=experiment_id,
                record_type="feed_step",
                stage=stage,
                section=section or "feed schedule",
                label=interval,
                source=source,
                occurrence=occurrence,
                planned_value=display_number(feed_rate),
                units="mL/min" if feed_rate is not None else "",
                elapsed_start_min=display_number(start_min),
                elapsed_end_min=display_number(end_min),
                feed_rate_mL_min=display_number(feed_rate),
                cumulative_percent=display_number(cumulative),
                radical_flux_mol_min_L=display_number(radical_flux),
                particle_size_nm=display_number(particle_size),
                source_range=row_source_range(source["source_range"], row_index),
                details_json=json.dumps(details, sort_keys=True, default=str),
            )
        )
        consumed.add(row_index)
    return records, consumed


def metric_records_from_row(
    row: list[Any],
    row_index: int,
    experiment_id: str,
    stage: str,
    section: str,
    source: dict[str, str],
    occurrence: dict[tuple[str, str, str, str], int],
) -> list[dict[str, Any]]:
    records = []
    for index, cell in enumerate(row):
        label = cell_text(cell)
        normalized = normalized_label(label)
        if not label or is_number(label) or len(label) > 160:
            continue
        metric = next(
            (
                (field, units, record_type)
                for terms, field, units, record_type in _IMPORTANT_METRICS
                if any(term in normalized for term in terms)
            ),
            None,
        )
        if metric is None:
            continue
        value = next_numeric_value(row, index + 1, index + 5)
        if value is None:
            continue
        field, units, record_type = metric
        value_field = (
            "planned_value"
            if any(
                term in normalized
                for term in ("target", "goal", "calculated", "estimated", "planned")
            )
            else "actual_value"
            if record_type == "result"
            else "planned_value"
        )
        kwargs: dict[str, Any] = {
            value_field: display_number(value),
            "units": units,
            "source_range": row_source_range(source["source_range"], row_index),
            "details_json": details_json_for_row(row),
        }
        if field:
            kwargs[field] = display_number(value)
        records.append(
            make_record(
                experiment_id=experiment_id,
                record_type=record_type,
                stage=stage,
                section=section,
                label=label,
                source=source,
                occurrence=occurrence,
                **kwargs,
            )
        )
    return records


def make_record(
    *,
    experiment_id: str,
    record_type: str,
    stage: str,
    section: str,
    label: str,
    source: dict[str, str],
    occurrence: dict[tuple[str, str, str, str], int],
    **values: Any,
) -> dict[str, Any]:
    semantic_key = (
        record_type,
        normalized_label(stage),
        normalized_label(section),
        normalized_label(label),
    )
    ordinal = occurrence.get(semantic_key, 0) + 1
    occurrence[semantic_key] = ordinal
    digest_input = "|".join(
        (
            source["source_key"],
            record_type,
            normalized_label(stage),
            normalized_label(section),
            normalized_label(label),
            str(ordinal),
        )
    )
    record = {
        "record_id": f"PN-{hashlib.sha256(digest_input.encode('utf-8')).hexdigest()[:16]}",
        "experiment_id": experiment_id,
        "source_key": source["source_key"],
        "record_type": record_type,
        "stage": stage,
        "section": section,
        "label": label,
        "source_fingerprint": source["source_fingerprint"],
        "synced_at": source["synced_at"],
        "active": "true",
        **values,
    }
    return project_row(PROJECT_RECORD_SHEET, record)


def infer_material_role(material: str, stage: str, section: str) -> str:
    text = normalized_label(f"{material} {section}")
    for terms, role in _ROLE_PATTERNS:
        if any(term in text for term in terms):
            return role
    if "monomer" in text or any(token in text for token in ("acrylate", "styrene", "vinyl", "ses", "ska")):
        if normalized_label(stage) == "core":
            return "core_monomer"
        if normalized_label(stage) in {"shell", "functional shell"}:
            return "shell_monomer"
        return "comonomer"
    return "additive"


def stage_from_row(row: list[Any]) -> str:
    nonempty = [cell_text(value) for value in row if cell_text(value)]
    if not nonempty or len(nonempty) > 5:
        return ""
    text = normalized_label(" ".join(nonempty))
    if "feed rate" in text or "time interval" in text:
        return ""
    for pattern, stage in _STAGE_PATTERNS:
        if re.search(rf"\b{re.escape(pattern)}(?:\s+stage)?\b", text):
            return stage
    return ""


def is_feed_header(row: list[Any]) -> bool:
    return any("time interval" in normalized_label(value) for value in row)


def note_from_row(row: list[Any]) -> str:
    text_cells = [
        cell_text(value)
        for value in row
        if cell_text(value) and not is_number(value)
    ]
    if not text_cells:
        return ""
    joined = " | ".join(text_cells)
    normalized = normalized_label(joined)
    if any(term in normalized for term in _NOTE_TERMS) and (
        len(joined) >= 24 or any(term in normalized for term in ("observ", "separat", "coag", "foam", "spray", "freeze"))
    ):
        return joined
    return ""


def nearest_section_label(values: list[list[Any]], row_index: int) -> str:
    for index in range(row_index - 1, max(-1, row_index - 8), -1):
        cells = [cell_text(value) for value in values[index] if cell_text(value)]
        if len(cells) == 1 and not is_number(cells[0]):
            return cells[0][:160]
    return ""


def find_source_date(values: list[list[Any]]) -> str:
    patterns = (
        r"\b(20\d{2})[-/](\d{1,2})[-/](\d{1,2})\b",
        r"\b(\d{1,2})/(\d{1,2})/(20\d{2})\b",
    )
    for row in values[:40]:
        for value in row:
            text = cell_text(value)
            match = re.search(patterns[0], text)
            if match:
                year, month, day = match.groups()
                return f"{int(year):04d}-{int(month):02d}-{int(day):02d}"
            match = re.search(patterns[1], text)
            if match:
                month, day, year = match.groups()
                return f"{int(year):04d}-{int(month):02d}-{int(day):02d}"
    return ""


def normalize_date_text(value: Any) -> str:
    text = cell_text(value)
    for pattern, order in (
        (r"\b(20\d{2})[-/](\d{1,2})[-/](\d{1,2})\b", "ymd"),
        (r"\b(\d{1,2})/(\d{1,2})/(20\d{2})\b", "mdy"),
    ):
        match = re.search(pattern, text)
        if not match:
            continue
        first, second, third = match.groups()
        if order == "ymd":
            year, month, day = first, second, third
        else:
            month, day, year = first, second, third
        return f"{int(year):04d}-{int(month):02d}-{int(day):02d}"
    return text


def lineage_experiment_id(description: str, product_lot: str) -> str:
    normalized_charge_id = charge_experiment_id(description, product_lot)
    if normalized_charge_id:
        return normalized_charge_id
    if product_lot and normalized_label(product_lot) not in {"n/a", "na"}:
        return re.sub(r"[^A-Za-z0-9._-]+", "-", product_lot).strip("-").upper()[:64]
    slug = re.sub(r"[^A-Za-z0-9]+", "-", description).strip("-").upper()
    return f"SRC-{slug[:56] or 'LINEAGE'}"


def experiment_id_from_sheet_name(sheet_name: str) -> str:
    match = re.search(r"\bEP\s*[- ]?\s*(\d+)(?:\s*-\s*([A-Za-z]+))?", sheet_name, re.IGNORECASE)
    if match:
        suffix = f"-{match.group(2).upper()}" if match.group(2) else ""
        return f"EP-{int(match.group(1)):03d}{suffix}"
    normalized_charge_id = charge_experiment_id(sheet_name)
    if normalized_charge_id:
        return normalized_charge_id
    slug = re.sub(r"[^A-Za-z0-9]+", "-", sheet_name).strip("-").upper()
    return f"SRC-{slug[:40] or 'NOTEBOOK'}"


def charge_experiment_id(description: str, product_lot: str = "") -> str:
    for value in (product_lot, description):
        match = re.search(r"\bBMS\s*[- ]\s*(\d+)", value, re.IGNORECASE)
        if match:
            return f"BMS-{int(match.group(1)):03d}"
    match = re.search(
        r"\b(?:\d+\s*L\s+)?Surfactant\s+Batch\s*#\s*(\d+)",
        description,
        re.IGNORECASE,
    )
    if match:
        return f"BMS-{int(match.group(1)):03d}"
    match = re.search(
        r"\bDES\s*-\s*(?:D?OX)\s+Surf\s*#\s*(\d+)",
        description,
        re.IGNORECASE,
    )
    if match:
        return f"DESOX-{int(match.group(1)):03d}"
    return ""


def build_source_key(
    spreadsheet_id: str,
    sheet_name: str,
    sheet_id: int | str | None,
) -> str:
    identity = str(sheet_id) if sheet_id not in (None, "") else sheet_name
    return f"{spreadsheet_id or 'local'}:{identity}"


def fingerprint_values(values: list[list[Any]]) -> str:
    payload = json.dumps(values, ensure_ascii=False, separators=(",", ":"), default=str)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def normalize_values(values: list[list[Any]]) -> list[list[Any]]:
    normalized = []
    for row in values:
        normalized.append(
            [
                value.strip() if isinstance(value, str) else ("" if value is None else value)
                for value in (row if isinstance(row, list) else [row])
            ]
        )
    while normalized and not any(cell_text(value) for value in normalized[-1]):
        normalized.pop()
    return normalized


def project_row(sheet_name: str, row: dict[str, Any]) -> dict[str, Any]:
    return {
        header: row.get(header, "")
        for header in sheet_by_name(sheet_name).headers
    }


def indexed_rows(
    rows: list[dict[str, Any]],
    key_field: str,
) -> dict[str, tuple[int, dict[str, Any]]]:
    indexed: dict[str, tuple[int, dict[str, Any]]] = {}
    for row_number, row in enumerate(rows, start=2):
        if not isinstance(row, dict):
            continue
        key_value = str(row.get(key_field, "")).strip()
        if not key_value:
            continue
        if key_value in indexed:
            first_row = indexed[key_value][0]
            raise ValueError(
                f"Duplicate existing {key_field} {key_value!r} at rows "
                f"{first_row} and {row_number}."
            )
        indexed[key_value] = (row_number, row)
    return indexed


def diff_row(
    sheet_name: str,
    row_number: int,
    key_field: str,
    key_value: str,
    current: dict[str, Any],
    desired: dict[str, Any],
) -> list[dict[str, Any]]:
    updates = []
    for field in sheet_by_name(sheet_name).headers:
        if field == key_field:
            continue
        if not same_cell(current.get(field), desired.get(field)):
            updates.append(
                cell_update(
                    sheet_name,
                    row_number,
                    key_field,
                    key_value,
                    field,
                    desired.get(field, ""),
                )
            )
    return updates


def cell_update(
    sheet_name: str,
    row_number: int,
    key_field: str,
    key_value: str,
    field: str,
    value: Any,
) -> dict[str, Any]:
    return {
        "sheet": sheet_name,
        "row_number": row_number,
        key_field: key_value,
        "key_field": key_field,
        "key_value": key_value,
        "field": field,
        "value": value,
    }


def same_cell(left: Any, right: Any) -> bool:
    return cell_text(left) == cell_text(right)


def deduplicate_records(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    unique: dict[str, dict[str, Any]] = {}
    for record in records:
        unique.setdefault(str(record.get("record_id", "")), record)
    return list(unique.values())


def details_json_for_row(row: list[Any]) -> str:
    return json.dumps(
        {f"column_{index + 1}": value for index, value in enumerate(row) if cell_text(value)},
        sort_keys=True,
        default=str,
    )


def row_source_range(source_range: str, zero_based_row_index: int) -> str:
    match = re.search(r"(\d+)", source_range)
    start_row = int(match.group(1)) if match else 1
    row_number = start_row + zero_based_row_index
    return f"{column_prefix(source_range)}{row_number}:{column_suffix(source_range)}{row_number}"


def column_prefix(source_range: str) -> str:
    match = re.match(r"\$?([A-Za-z]+)", source_range)
    return match.group(1).upper() if match else "A"


def column_suffix(source_range: str) -> str:
    match = re.search(r":\$?([A-Za-z]+)", source_range)
    return match.group(1).upper() if match else column_prefix(source_range)


def parse_time_interval(value: str) -> tuple[float | None, float | None]:
    numbers = [float(item) for item in re.findall(r"-?\d+(?:\.\d+)?", value)]
    if len(numbers) >= 2:
        return numbers[0], abs(numbers[1])
    if numbers:
        return numbers[0], None
    return None, None


def duration_minutes(value: Any) -> float | None:
    text = cell_text(value)
    match = re.fullmatch(r"(\d+):(\d{1,2}):(\d{1,2}(?:\.\d+)?)", text)
    if match:
        hours, minutes, seconds = match.groups()
        return int(hours) * 60 + int(minutes) + float(seconds) / 60
    return numeric_value(value)


def first_numeric_by_header(
    details: dict[str, Any],
    include: tuple[str, ...],
    exclude: tuple[str, ...] = (),
) -> float | None:
    for header, value in details.items():
        normalized = normalized_label(header)
        if all(term in normalized for term in include) and not any(term in normalized for term in exclude):
            parsed = numeric_value(value)
            if parsed is not None:
                return parsed
    return None


def next_numeric_value(row: list[Any], start: int, end: int) -> float | None:
    for index in range(start, min(end, len(row))):
        parsed = numeric_value(row[index])
        if parsed is not None:
            return parsed
    return None


def numeric_value(value: Any) -> float | None:
    if value is None or value == "" or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = cell_text(value).replace(",", "").strip()
    if not text:
        return None
    match = re.fullmatch(r"[-+]?\d+(?:\.\d+)?(?:[Ee][-+]?\d+)?\s*%?", text)
    if not match:
        return None
    try:
        return float(text.rstrip("%").strip())
    except ValueError:
        return None


def is_number(value: Any) -> bool:
    return numeric_value(value) is not None


def display_number(value: float | None) -> str:
    if value is None:
        return ""
    if float(value).is_integer():
        return str(int(value))
    return format(value, ".12g")


def first_index(values: list[Any], predicate: Any) -> int | None:
    for index, value in enumerate(values):
        if predicate(value):
            return index
    return None


def value_at(row: list[Any], index: int | None) -> Any:
    if index is None or index < 0 or index >= len(row):
        return ""
    return row[index]


def normalized_label(value: Any) -> str:
    return re.sub(r"\s+", " ", cell_text(value).lower()).strip()


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()
